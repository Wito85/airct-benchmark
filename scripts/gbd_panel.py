#!/usr/bin/env python3
"""GBD panel of the AI-RCT Benchmark (Protocol Amendment v3.0, Section 5, decision D24).

Four steps, each one command, run from the repository root:

  1. rank       ranks the GBD 2023 Level 3 causes by global DALYs 2023 from the export of the
                GBD Results Tool and applies the exclusion rules (a) to (c) of Section 5.2
  2. sheets     writes one mapping sheet per rater (Excel), pre-filled with the proposed MeSH
                expression and its facts from data/gbd/gbd_mesh_candidates.csv
  3. reconcile  checks the two filled sheets against the ranking and the candidates table,
                lists every disagreement, and, once the raters agree, writes the mapping table
                and the final panel
  4. build      rebuilds config/gbd_panel.yaml and the table for the amendment from a mapping
                table (called by reconcile; a final panel needs a reconciled mapping table)

Exit codes: 0 done; 1 the raters disagree and have to decide (nothing was written but the
consensus sheet); 2 stopped, the message says why.

Nothing in this script contacts PubMed or any other server. The raw IHME files and the ranking
with DALY values stay with the authors, in a folder outside the repository (the folder of the
raw export); they are not committed, not uploaded and not redistributed. Only their SHA-256
hashes, the export settings and the derived ranking without values are deposited.

Dependencies: Python 3.10 or newer, pyyaml (already required by the pipeline), openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import io
import json
import math
import os
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from airct_benchmark import gbd as G  # noqa: E402  (standard library only; no network)

GBD_DIR = REPO / "data" / "gbd"
CANDIDATES = GBD_DIR / "gbd_mesh_candidates.csv"
RANKING_FULL_NAME = "gbd2023_level3_dalys_global_2023.csv"   # with DALY values: written next to the raw export, never deposited
RANKING_PUBLIC = GBD_DIR / "gbd2023_level3_ranking.csv"
EXPORT_RECORD = GBD_DIR / "gbd_export_record.json"
MAPPING = GBD_DIR / "gbd_panel_mesh_mapping.csv"
AMENDMENT_TABLE = GBD_DIR / "gbd_panel_amendment_table.md"
RATING_DIR = GBD_DIR / "rating"
PANEL_YAML = REPO / "config" / "gbd_panel.yaml"
CATALOGUE_YAML = REPO / "config" / "catalogue.yaml"

YEAR = 2023
PANEL_SIZE = G.PANEL_SIZE
DEFAULT_ROWS = 30          # causes per rater sheet: 20 for the panel and 10 in reserve
ALL_CAUSES_TOLERANCE = 0.005
# The stratum of Amendment v3.0, Section 5.1, as the GBD Results Tool identifies it: (label, id column, id, name columns,
# English names). The ids do not depend on the language in which the tool was used; the English names serve an
# export that carries no ids. Measure 2 = DALYs, location 1 = Global, sex 3 = Both, age 22 = All ages, metric 1 = Number.
STRATUM = (
    ("measure", "measure_id", "2", ("measure_name", "measure"), ("dalys (disability-adjusted life years)", "dalys")),
    ("location", "location_id", "1", ("location_name", "location"), ("global",)),
    ("sex", "sex_id", "3", ("sex_name", "sex"), ("both", "both sexes")),
    ("age", "age_id", "22", ("age_name", "age"), ("all ages", "all age")),
    ("metric", "metric_id", "1", ("metric_name", "metric"), ("number",)),
)
POPULATION_GROUP_ALL = "1"       # column population_group_id of the GBD 2023 Results Tool: 1 = all population
ALL_CAUSES_ID = "294"
NAME_SIMILARITY = 0.75      # same cause id in export and hierarchy: below this the two names are not one name spelled differently
IHME_SOURCE_LINE = "Source: Institute for Health Metrics and Evaluation. Used with permission. All rights reserved."
GBD_CITATION = ("GBD 2023 Disease and Injury and Risk Factor Collaborators. Lancet 2025;406(10513):1873 to 1922, "
                "DOI 10.1016/S0140-6736(25)01637-X")
MESH_BROWSER = "https://meshb.nlm.nih.gov/record/ui?ui="

EXPRESSION_RE = G.EXPRESSION_RE
DECISIONS = ("agree", "use alternative", "no single descriptor (rule d)")
NONE_PROPOSED = "(none proposed: rule d)"
CONSENSUS_KEY = "CONSENSUS"
SHEET_COLUMNS = ["rank", "gbd_cause", "proposed_expression", "mesh_ui", "mesh_browser_link", "date_introduced", "scope_relation",
                 "mesh_scope_note", "narrower_descriptors", "inclusion_note", "alternatives",
                 "DECISION", "ALTERNATIVE_EXPRESSION", "COMMENT"]
SHEET_REQUIRED = ["rank", "gbd_cause", "proposed_expression", "DECISION", "ALTERNATIVE_EXPRESSION", "COMMENT"]
FACT_COLUMNS = ["mesh_ui", "date_introduced", "scope_relation", "inclusion_note"]      # facts that enter the panel
CANDIDATE_COLUMNS = ["gbd_cause", "role", "expression", "mesh_ui", "date_introduced", "scope_relation", "mesh_scope_note",
                     "narrower_descriptors", "inclusion_note", "alternatives"]
CANDIDATE_ROLES = ("proposed", "registered", "alternative", "no single descriptor")
SCOPE_RELATIONS = ("exact", "broader", "narrower", "mixed", "none")


class GbdError(Exception):
    """A problem the user has to resolve; printed without a traceback."""


# ---------------------------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------------------------

norm_name = G.norm_name
gbd_key = G.gbd_key
sha256_text = G.sha256_text


def norm_col(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def sha256_file(path: Path) -> str:
    """SHA-256 of the bytes of a file. Used for the raw IHME files, which are never rewritten."""
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def sha256_of_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def write_text_lf(path: Path, text: str) -> None:
    """Write UTF-8 text with LF line endings on every platform."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(text)


def csv_text(columns: list[str], rows: list[dict]) -> str:
    buf = io.StringIO(newline="")
    w = csv.DictWriter(buf, fieldnames=columns, lineterminator="\n", extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({c: r.get(c, "") for c in columns})
    return buf.getvalue()


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    write_text_lf(path, csv_text(columns, rows))


def read_table(path: Path, what: str, needed: list[str]) -> list[dict]:
    """Rows of a CSV file written by this script (or the candidates table), with a column check."""
    if not path.is_file():
        raise GbdError(f"{what} not found: {shown(path)}")
    try:
        rows = G.read_csv_dicts(path)
    except UnicodeDecodeError as exc:
        raise GbdError(f"{what} {shown(path)} is not UTF-8 text; it was probably saved by a spreadsheet program. Restore the file.") from exc
    if not rows:
        raise GbdError(f"{what} {shown(path)} has no data rows")
    missing = [c for c in needed if c not in rows[0]]
    if missing:
        raise GbdError(f"{what} {shown(path)} lacks the column(s) {missing}; it is not the file this step expects")
    return rows


def rel(path: Path, root: Path | None = None) -> str:
    """Path relative to the repository (or to ``root``); a file outside it is named without its folder."""
    root = root or REPO
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        return path.name


def shown(path: Path) -> str:
    """Path as printed for the user: relative inside the repository, complete outside it."""
    try:
        return path.resolve().relative_to(REPO.resolve()).as_posix()
    except ValueError:
        return str(path)


# ---------------------------------------------------------------------------------------------
# 1. rank
# ---------------------------------------------------------------------------------------------

def _pick(row: dict, *names: str) -> str:
    for n in names:
        if n in row and str(row[n]).strip() != "":
            return str(row[n]).strip()
    return ""


def read_export(path: Path) -> tuple[list[dict], list[str]]:
    """The CSV of the GBD Results Tool exactly as downloaded -> (rows with normalised column names, notes)."""
    raw = path.read_bytes()
    notes: list[str] = []
    if raw[:2] == b"PK":
        raise GbdError(f"{path.name} is a ZIP archive (or an Excel file), not a CSV file. The Results Tool delivers a ZIP: "
                       "unpack it and pass the CSV file that is inside.")
    if raw[:3] == b"\xef\xbb\xbf":
        notes.append("the export begins with a byte order mark")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise GbdError(f"{path.name} is not UTF-8 text. Pass the CSV file exactly as downloaded from the Results Tool.") from exc
    first = text.split("\n", 1)[0]
    if first.count(";") > first.count(","):
        raise GbdError(f"{path.name} is separated by semicolons. The download of the Results Tool is separated by commas, so this "
                       "file was opened and saved again (for example in Excel). Use the untouched download.")
    rows = [{norm_col(k or ""): ("" if v is None else str(v).strip()) for k, v in r.items()}
            for r in csv.DictReader(io.StringIO(text, newline=""))]
    if not rows:
        raise GbdError(f"{path.name}: the file has no data rows")
    return rows, notes


def read_hierarchy(path: Path) -> tuple[dict[str, dict], str]:
    """GBD cause hierarchy (xlsx or csv) -> ({cause_id: {name, parent_id, level}}, where it was read from)."""
    rows: list[dict] = []
    origin = path.name
    needed = ("cause_id", "cause_name", "parent_id", "level")
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise GbdError(f"{path.name} cannot be read as an Excel workbook (.xlsx): {exc}. Download the hierarchy file again.") from exc
        try:
            for ws in wb.worksheets:
                try:
                    ws.reset_dimensions()      # some workbooks state a wrong sheet size
                except AttributeError:
                    pass
                header = None
                for values in ws.iter_rows(values_only=True):
                    cells = ["" if v is None else str(v).strip() for v in values]
                    if header is None:
                        cols = [norm_col(c) for c in cells]
                        if all(c in cols for c in needed):
                            header = cols
                        continue
                    if any(cells):
                        rows.append(dict(zip(header, cells)))
                if rows:
                    origin = f"{path.name}, worksheet '{ws.title}'"
                    break
        finally:
            wb.close()                         # a workbook opened read-only keeps the file open (it would stay locked on Windows)
    elif path.suffix.lower() == ".xls":
        raise GbdError(f"{path.name} is an old Excel file (.xls). Download the hierarchy as .xlsx or .csv.")
    else:
        try:
            rows = [{norm_col(k): v for k, v in r.items()} for r in G.read_csv_dicts(path)]
        except UnicodeDecodeError as exc:
            raise GbdError(f"{path.name} is not UTF-8 text; pass the hierarchy file as downloaded (.xlsx or .csv)") from exc
        if rows and not all(c in rows[0] for c in needed):
            rows = []
    if not rows:
        raise GbdError(f"{path.name}: no table with the columns 'Cause ID', 'Cause Name', 'Parent ID' and 'Level' was found. "
                       "This must be the GBD 2023 cause hierarchy file from GHDx.")
    out: dict[str, dict] = {}
    for r in rows:
        cid = _pick(r, "cause_id").split(".")[0]
        if not cid:
            continue
        try:
            level = int(float(_pick(r, "level")))
        except ValueError:
            continue
        out[cid] = {"name": _pick(r, "cause_name"), "parent_id": _pick(r, "parent_id").split(".")[0], "level": level}
    if not any(v["level"] == 3 for v in out.values()):
        raise GbdError(f"{path.name}: the hierarchy has no Level 3 causes; wrong file?")
    return out, origin


def level1_group(cid: str, hierarchy: dict[str, dict]) -> str:
    seen = set()
    cur = cid
    while cur in hierarchy and hierarchy[cur]["level"] > 1 and cur not in seen:
        seen.add(cur)
        cur = hierarchy[cur]["parent_id"]
    node = hierarchy.get(cur)
    if node is None or node["level"] != 1:
        return "unresolved"
    low = node["name"].lower()
    if low.startswith("non-communicable"):
        return "NCD"
    if low.startswith("communicable"):
        return "CMNN"
    if low.startswith("injur"):
        return "Injuries"
    return f"unresolved: {node['name']}"


def cmd_rank(args: argparse.Namespace) -> int:
    export, hier_path = Path(args.export), Path(args.hierarchy)
    for p in (export, hier_path):
        if not p.is_file():
            raise GbdError(f"file not found: {p}")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", args.export_date):
        raise GbdError(f"--export-date must be written as YYYY-MM-DD, for example 2026-10-07 (found {args.export_date!r})")
    if not args.author.strip() or not args.permalink.strip().lower().startswith("http"):
        raise GbdError("--author needs a name and --permalink the web address of the Results Tool query (it begins with http)")

    citation = ""
    if args.citation_file:
        if not Path(args.citation_file).is_file():
            raise GbdError(f"file not found: {args.citation_file}")
        citation = G.one_line(Path(args.citation_file).read_text(encoding="utf-8-sig", errors="replace"))

    rows, notes = read_export(export)
    if any(c.startswith("rei") for c in rows[0]):
        raise GbdError(f"{export.name} has risk-factor columns (rei). In the Results Tool, 'GBD Estimate' must be "
                       "'Cause of death or injury'. Export again.")
    columns = set(rows[0])
    for label, id_col, _, name_cols, _ in STRATUM:
        if id_col not in columns and not any(n in columns for n in name_cols):
            raise GbdError(f"{export.name}: column for '{label}' not found (looked for {[id_col, *name_cols]}; found {sorted(columns)}). "
                           "Download the CSV from the GBD Results Tool with names and ids (option 'Both').")
    for label, names in {"cause": ("cause_id", "cause_name", "cause"), "year": ("year", "year_id"), "val": ("val", "value", "mean")}.items():
        if not any(n in columns for n in names):
            raise GbdError(f"{export.name}: column for '{label}' not found (looked for {list(names)}; found {sorted(columns)}). "
                           "Download the CSV from the GBD Results Tool with names and ids (option 'Both').")

    def ident(r: dict, column: str) -> str:
        return _pick(r, column).split(".")[0]

    def in_dimension(r: dict, id_col: str, wanted: str, name_cols: tuple[str, ...], english: tuple[str, ...]) -> bool:
        """By id where the export carries one (the same in every language of the tool), else by the English name."""
        if ident(r, id_col):
            return ident(r, id_col) == wanted
        return _pick(r, *name_cols).lower() in english

    def keep(r: dict) -> bool:
        if ident(r, "population_group_id") not in ("", POPULATION_GROUP_ALL):
            return False
        return (all(in_dimension(r, id_col, wanted, name_cols, english) for _, id_col, wanted, name_cols, english in STRATUM)
                and _pick(r, "year", "year_id").split(".")[0] == str(YEAR))

    sel = [r for r in rows if keep(r)]
    if not sel:
        # Say which setting is wrong: the dimension for which no row of the file has the value of the protocol.
        wanted_names = {"measure": "DALYs", "location": "Global", "sex": "Both", "age": "All ages", "metric": "Number"}
        wrong = []
        for label, id_col, wanted, name_cols, english in STRATUM:
            if not any(in_dimension(r, id_col, wanted, name_cols, english) for r in rows):
                held = sorted({" / ".join(x for x in (ident(r, id_col), _pick(r, *name_cols)) if x) for r in rows})[:6]
                wrong.append(f"{label} must be {wanted_names[label]} (id {wanted}), the file holds {held}")
        if not any(_pick(r, "year", "year_id").split(".")[0] == str(YEAR) for r in rows):
            wrong.append(f"year must be {YEAR}, the file holds {sorted({_pick(r, 'year', 'year_id') for r in rows})[:6]}")
        if not any(ident(r, "population_group_id") in ("", POPULATION_GROUP_ALL) for r in rows):
            wrong.append(f"population group must be the whole population (id {POPULATION_GROUP_ALL})")
        found = {label: sorted({" / ".join(x for x in (ident(r, id_col), _pick(r, *name_cols)) if x) for r in rows})[:8]
                 for label, id_col, _, name_cols, _ in STRATUM}
        found["year"] = sorted({_pick(r, "year", "year_id") for r in rows})[:8]
        raise GbdError(f"{export.name}: no row matches Measure DALYs, Location Global, Sex Both, Age All ages, Metric Number, "
                       f"Year {YEAR}. " + ("Wrong setting: " + "; ".join(wrong) + ". Export again with this setting changed. " if wrong else
                                           f"Values found in the file: {json.dumps(found, ensure_ascii=False)}. ")
                       + "An export without ids is read by its English names only: download it with names and ids (option 'Both').")
    # The names of the stratum as the export writes them (the Results Tool can be used in several languages).
    stratum_named = {label: _pick(sel[0], *name_cols) for label, _, _, name_cols, _ in STRATUM}
    not_english = {label: name for (label, _, _, _, english), name in zip(STRATUM, stratum_named.values())
                   if name and name.lower() not in english}
    if not_english:
        notes.append("the export is not in English (" + ", ".join(f"{k} '{v}'" for k, v in not_english.items()) + "): the stratum "
                     "and the causes were identified by their ids, and the cause names are those of the hierarchy file")

    hierarchy, hierarchy_origin = read_hierarchy(hier_path)
    ids_by_name: dict[str, list[str]] = {}
    for cid, v in hierarchy.items():
        ids_by_name.setdefault(norm_name(v["name"]), []).append(cid)
    causes: dict[str, dict] = {}
    unknown: list[str] = []
    name_differences: list[str] = []
    all_causes_value: float | None = None

    def number(text: str, name: str) -> float:
        try:
            val = float(text)
        except ValueError as exc:
            hint = " (decimal comma: the file was saved again by a spreadsheet program; use the untouched download)" if "," in text else ""
            raise GbdError(f"{export.name}: value of '{name}' is not a number: {text!r}{hint}") from exc
        if not math.isfinite(val) or val < 0:
            raise GbdError(f"{export.name}: value of '{name}' is not a finite, non-negative number: {text!r}")
        return val

    renamed = 0
    for r in sel:
        name = _pick(r, "cause_name", "cause")
        cid = ident(r, "cause_id")
        if cid not in hierarchy:
            # export without cause ids: match by name, which must then be unambiguous in the hierarchy
            same_name = ids_by_name.get(norm_name(name), [])
            if len(same_name) > 1:
                raise GbdError(f"{export.name}: the cause name '{name}' occurs {len(same_name)} times in the hierarchy and the export "
                               "has no cause ids. Download the export again with names and ids (option 'Both').")
            cid = same_name[0] if same_name else ""
        if cid not in hierarchy:
            if ident(r, "cause_id") == ALL_CAUSES_ID or norm_name(name) == "allcauses":      # a hierarchy file without the row of Level 0
                all_causes_value = number(_pick(r, "val", "value", "mean"), name)
            else:
                unknown.append(name or f"cause id {ident(r, 'cause_id')}")
            continue
        official = hierarchy[cid]["name"]
        if name and norm_name(official) != norm_name(name):
            renamed += 1
            if not not_english:
                # Same cause id, another English name. A different spelling of the same name is noted; a different
                # name means that export and hierarchy are not the same GBD round.
                similarity = difflib.SequenceMatcher(None, norm_name(official), norm_name(name)).ratio()
                if similarity < NAME_SIMILARITY:
                    raise GbdError(f"{export.name}: cause id {cid} is '{name}' in the export but '{official}' in the hierarchy "
                                   f"file; export and hierarchy must both be GBD {YEAR}. If both files are GBD {YEAR}, report this "
                                   "message; nothing was written.")
                name_differences.append(f"cause id {cid}: '{name}' in the export, '{official}' in the hierarchy file")
        if cid in causes:
            raise GbdError(f"{export.name}: cause '{official}' occurs twice after filtering; the export must hold one row per cause. "
                           "If the download came in several CSV files, the query was too large: select only Location Global and Year 2023.")
        # The name of a cause is the name of the hierarchy file; the export's own name is kept for the rule check below.
        causes[cid] = {"cause_id": cid, "cause_name": official, "export_name": name or official, "level": hierarchy[cid]["level"],
                       "level1_group": level1_group(cid, hierarchy), "daly_count": number(_pick(r, "val", "value", "mean"), official),
                       "daly_lower": _pick(r, "lower"), "daly_upper": _pick(r, "upper")}
    if not causes:
        raise GbdError(f"{export.name}: none of its {len(sel)} cause(s) was found in the hierarchy file. The export has no cause ids "
                       "and is not in English, or the hierarchy file is not the GBD cause hierarchy. Download the export with names "
                       "and ids (option 'Both').")
    if not_english and renamed:
        notes.append(f"{renamed} cause name(s) of the export differ from the names of the hierarchy file (language of the export)")

    # What is ranked. The Level 3 causes of the hierarchy, and, should the hierarchy hold one, a cause of Level 1
    # or 2 that has no cause below it: GBD reports such a cause at Level 3 as itself, and without it the ranked
    # causes would not add up to all causes. Aggregates across causes (names that begin with "Total") are no part
    # of the partition, wherever they stand; they are never ranked and are recorded.
    def is_aggregate(name: str) -> bool:
        return name.strip().lower().startswith("total ")

    parents = {v["parent_id"] for cid, v in hierarchy.items() if v["parent_id"] != cid}
    level3_ids = {cid for cid, v in hierarchy.items() if v["level"] == 3 and not is_aggregate(v["name"])}
    leaf_ids = {cid for cid, v in hierarchy.items() if v["level"] in (1, 2) and cid not in parents and not is_aggregate(v["name"])}
    hierarchy_aggregates = sorted(v["name"] for v in hierarchy.values() if v["level"] in (1, 2, 3) and is_aggregate(v["name"]))
    # The GBD 2023 hierarchy file lists five such aggregates at Level 1 ("Total cancers" and others); they are never
    # ranked. Only an aggregate listed at Level 3 could change a ranking, so only that needs the total to vouch for it.
    level3_aggregates = sorted(v["name"] for v in hierarchy.values() if v["level"] == 3 and is_aggregate(v["name"]))
    export_aggregates = sorted(n for n in unknown if is_aggregate(n))
    foreign = sorted(n for n in unknown if not is_aggregate(n))

    missing = sorted(hierarchy[c]["name"] for c in level3_ids - set(causes))
    if missing and not (set(causes) & level3_ids):
        raise GbdError(f"{export.name} holds no Level 3 cause at all ({len(sel)} row(s) of the stratum: "
                       f"{sorted(c['cause_name'] for c in causes.values())[:5]}). In the Results Tool the field 'Cause' was left at "
                       f"its default. Select all causes (every level, {len(hierarchy)} entries) and export again; the CSV then holds "
                       "one row per cause.")
    if missing:
        raise GbdError(f"{len(missing)} Level 3 cause(s) of the hierarchy are missing from the export, so ranks would be wrong: "
                       f"{missing[:12]}{' ...' if len(missing) > 12 else ''}. In the Results Tool select all causes "
                       "(or all Level 3 causes) and export again.")
    leaves_without_value = sorted(hierarchy[c]["name"] for c in leaf_ids - set(causes))
    ranked = sorted((c for cid, c in causes.items() if cid in level3_ids or cid in leaf_ids), key=lambda c: (-c["daly_count"], c["cause_name"]))
    for c in ranked:
        c["above_level3"] = c["cause_id"] in leaf_ids
    unresolved = [f"{c['cause_name']} ({c['level1_group']})" for c in ranked if c["level1_group"] not in ("NCD", "CMNN", "Injuries")]
    if unresolved:
        raise GbdError("the Level 1 group (non-communicable; communicable, maternal, neonatal and nutritional; injuries) could not be "
                       f"read from the hierarchy for {len(unresolved)} cause(s), so rule (a) cannot be applied: {unresolved[:6]}. "
                       "The hierarchy file needs the columns Parent ID and Level for every cause up to Level 1.")
    for i, c in enumerate(ranked, 1):
        c["rank"] = i
    tie_pairs = [(a, b) for a, b in zip(ranked, ranked[1:]) if a["daly_count"] == b["daly_count"]]
    top_ties = [(a["cause_name"], b["cause_name"]) for a, b in tie_pairs if a["rank"] <= 60]
    if top_ties:
        raise GbdError(f"two causes among the first 60 have exactly the same value, so their order is undefined: {top_ties[:3]}. "
                       "This does not happen with a genuine export; check the file.")
    if tie_pairs:
        notes.append(f"{len(tie_pairs)} pair(s) of causes below rank 60 have identical values and are ordered by name")

    # Completeness check: the ranked causes partition all causes, so their DALY counts add up to 'All causes'.
    # It is what guards the ranking against a hierarchy of another GBD round, a missing cause or a cause counted twice.
    level0 = [c for c in causes.values() if c["level"] == 0]
    if level0:
        all_causes_value = level0[0]["daly_count"]
    ranked_sum = sum(c["daly_count"] for c in ranked)
    if all_causes_value:
        ratio = ranked_sum / all_causes_value
        if abs(ratio - 1.0) > ALL_CAUSES_TOLERANCE:
            raise GbdError(f"the DALY counts of the {len(ranked)} ranked causes add up to {ratio:.4f} times the value of "
                           f"'All causes'. They must add up to it (tolerance {ALL_CAUSES_TOLERANCE:.1%}); "
                           "export or hierarchy are incomplete or from different GBD rounds."
                           + (f" Causes of the export that are not in the hierarchy file: {foreign[:10]}." if foreign else "")
                           + (f" Causes without a value: {leaves_without_value[:10]}." if leaves_without_value else ""))
        all_causes_check = {"done": True, "level3_sum_over_all_causes": round(ratio, 6), "tolerance": ALL_CAUSES_TOLERANCE}
    else:
        all_causes_check = {"done": False, "reason": "the export has no row for 'All causes' (Level 0)"}
        notes.append("the sum of the ranked causes could not be compared with 'All causes' (no such row in the export)")

    # Without a passed completeness check nothing unexpected is tolerated. With it, the following are reported
    # and recorded, because they cannot change the order of the leading causes.
    unexpected = []
    if foreign:
        unexpected.append(f"{len(foreign)} cause(s) of the export are not in the hierarchy file: {foreign[:10]}{' ...' if len(foreign) > 10 else ''}")
    if level3_aggregates:
        unexpected.append(f"the hierarchy file lists aggregates across causes among the causes of Level 3: {level3_aggregates}")
    if leaves_without_value:
        unexpected.append(f"cause(s) of Level 1 or 2 without causes below them have no row in the export: {leaves_without_value}")
    if unexpected and not all_causes_check["done"]:
        raise GbdError("; ".join(unexpected) + f". This could not be checked against the total, because the export has no row for "
                       f"'All causes'. In the Results Tool select all causes (every level) and export again; export and hierarchy "
                       f"must both be GBD {YEAR}.")
    if foreign:
        notes.append(f"{len(foreign)} cause(s) of the export are not in the hierarchy file and were not ranked (the ranked causes add up "
                     f"to 'All causes', so no leading cause is missing): {foreign}")
    if hierarchy_aggregates:
        notes.append(f"aggregates across causes listed in the hierarchy file at Level 1 to 3 were not ranked: {hierarchy_aggregates}")
    if leaves_without_value:
        notes.append(f"cause(s) of Level 1 or 2 without causes below them have no row in the export and were not ranked: {leaves_without_value}")
    for d in name_differences:
        notes.append(f"spelling differs, treated as the same cause: {d}")

    def rule(name: str, group: str) -> tuple[str, str]:
        if group == "Injuries":
            return "excluded", "rule (a): Level 1 group Injuries"
        if "pandemic-related" in name.lower():
            return "excluded", "rule (b): pandemic-related outcome category"
        if name.strip().lower().startswith("other "):
            return "excluded", "rule (c): residual cause (name begins with Other)"
        return G.KEPT, ""

    for c in ranked:
        c["panel_rule_result"], c["exclusion_reason"] = rule(c["cause_name"], c["level1_group"])
        if not not_english and rule(c["export_name"], c["level1_group"])[0] != c["panel_rule_result"]:
            raise GbdError(f"cause id {c['cause_id']} is '{c['export_name']}' in the export and '{c['cause_name']}' in the hierarchy "
                           "file, and the exclusion rules (b) and (c), which go by the name, give different results for the two "
                           "names. Report this message; nothing was written.")
        if c["above_level3"]:
            if c["panel_rule_result"] == G.KEPT:
                raise GbdError(f"'{c['cause_name']}' is a cause of Level {c['level']} without Level 3 causes below it and is not "
                               "excluded by rule (a), (b) or (c). The amendment defines the panel on Level 3 causes and does not say "
                               "how such a cause is treated. Report this message; the rule has to be stated in the amendment before "
                               "the ranking is made. Nothing was written.")
            c["exclusion_reason"] += f"; cause of Level {c['level']} without Level 3 causes below it, ranked as itself"
    above = [f"{c['cause_name']} (Level {c['level']}, rank {c['rank']}, {c['exclusion_reason']})" for c in ranked if c["above_level3"]]
    if above:
        notes.append(f"ranked with the Level 3 causes, as GBD reports them: {above}")
    level3 = ranked

    # The ranking with DALY values is written next to the raw export, that is outside the repository
    # when the raw files are kept outside it, as they should be (IHME user agreement: no redistribution).
    local_dir = Path(args.local_dir) if args.local_dir else export.resolve().parent
    ranking_full = local_dir / RANKING_FULL_NAME
    write_csv(ranking_full, ["rank", "cause_id", "cause_name", "level1_group", "daly_count", "daly_lower", "daly_upper",
                             "panel_rule_result", "exclusion_reason"], level3)
    write_csv(RANKING_PUBLIC, ["rank", "cause_id", "cause_name", "level1_group", "panel_rule_result", "exclusion_reason"], level3)
    record = {
        "source": f"Global Burden of Disease Study 2023 (GBD 2023), {GBD_CITATION}",
        "source_identifier": IHME_SOURCE_LINE,
        "data_path": "GBD Results Tool, https://vizhub.healthdata.org/gbd-results/",
        "settings": {"measure": "DALYs (Disability-Adjusted Life Years)", "metric": "Number", "location": "Global",
                     "age": "All ages", "sex": "Both", "year": YEAR, "cause": "all causes or all Level 3 causes"},
        "stratum_as_named_in_the_export": stratum_named,
        "stratum_identified_by": "ids of the GBD Results Tool (measure 2, location 1, sex 3, age 22, metric 1, population group 1) "
                                 "where the export carries them, else English names; cause names from the hierarchy file",
        "data_citation_delivered_with_the_download": citation,
        "permalink": args.permalink.strip(),
        "exported_by": args.author.strip(),
        "export_date": args.export_date,
        "ranked_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "raw_export": {"file_name": export.name, "sha256": sha256_file(export), "rows": len(rows)},
        "cause_hierarchy": {"file_name": hier_path.name, "sha256": sha256_file(hier_path), "read_from": hierarchy_origin},
        "level3_causes_ranked": len(level3),
        "all_causes_check": all_causes_check,
        "export_aggregates_outside_hierarchy_ignored": export_aggregates,
        "export_causes_outside_hierarchy_not_ranked": foreign,
        "hierarchy_aggregates_not_ranked": hierarchy_aggregates,
        "causes_above_level_3_ranked_as_themselves": [c["cause_name"] for c in level3 if c["above_level3"]],
        "names_spelled_differently_in_export_and_hierarchy": name_differences,
        "notes": notes,
        "ranking_with_values_kept_by_the_authors": {"file_name": ranking_full.name, "sha256": sha256_file(ranking_full)},
        "ranking_public": {"file": rel(RANKING_PUBLIC), "sha256": sha256_text(RANKING_PUBLIC)},
        "hash_rule": "raw IHME files: SHA-256 of the file bytes; text files written by this script: SHA-256 of the "
                     "UTF-8 content with LF line endings",
        "note": "The raw export, the hierarchy file and the ranking with DALY values are kept by the authors and are "
                "not redistributed; their SHA-256 hashes are stated here.",
    }
    write_text_lf(EXPORT_RECORD, json.dumps(record, indent=2, ensure_ascii=False) + "\n")

    kept = [c for c in level3 if c["panel_rule_result"] == G.KEPT]
    n_above = sum(1 for c in level3 if c["above_level3"])
    print(f"{len(level3)} causes ranked by global DALYs {YEAR} ({len(level3) - n_above} of Level 3"
          + (f", {n_above} of Level 1 or 2 without causes below them" if n_above else "") + f"); {len(kept)} kept after rules (a) to (c).")
    print(f"hierarchy read from: {hierarchy_origin}")
    if all_causes_check["done"]:
        print(f"completeness: the ranked causes add up to {all_causes_check['level3_sum_over_all_causes']:.4f} times 'All causes' (expected 1)")
    for n in notes:
        print(f"NOTE: {n}")
    if export_aggregates:
        print(f"NOTE: {len(export_aggregates)} aggregate(s) across causes are in the export, are no causes of the hierarchy and were "
              f"ignored: {export_aggregates}")
    print(f"\n{'rank':>4}  {'group':<8}  {'DALYs (million)':>15}  cause")
    count = 0
    for c in level3:
        mark = "" if c["panel_rule_result"] == G.KEPT else f"   <- excluded, {c['exclusion_reason']}"
        print(f"{c['rank']:>4}  {c['level1_group']:<8}  {c['daly_count'] / 1e6:>15.1f}  {c['cause_name']}{mark}")
        count += c["panel_rule_result"] == G.KEPT
        if count >= args.show:
            break
    print(f"\nwritten: {ranking_full} (with DALY values; stays with you, never upload or pass on)")
    print(f"written: {shown(RANKING_PUBLIC)} (ranks and rule results, for the repository)")
    print(f"written: {shown(EXPORT_RECORD)} (hashes and export settings)")
    inside = [p for p in (export.resolve(), hier_path.resolve(), ranking_full.resolve()) if REPO.resolve() in p.parents]
    if inside:
        print("\nNOTE: these files lie inside the repository folder. Move them to a folder outside it before you upload or "
              "commit anything; they must not be published (IHME user agreement):")
        for path in inside:
            print(f"  {shown(path)}")
    existing = sorted(RATING_DIR.glob("gbd_mapping_rater_*.xlsx")) if RATING_DIR.is_dir() else []
    if existing:
        print("\nRater sheets exist already. They stay valid as long as the ranking did not change; 'reconcile' checks this.")
        print("Next step: python scripts/gbd_panel.py reconcile")
    else:
        print("\nNext step: python scripts/gbd_panel.py sheets")
    return 0


# ---------------------------------------------------------------------------------------------
# candidates
# ---------------------------------------------------------------------------------------------

@dataclass
class Candidates:
    rows: list[dict]
    path: Path

    def proposal(self, cause: str) -> dict | None:
        """The row shown to the raters for a cause: proposed, registered (named field) or 'no single descriptor'."""
        key = norm_name(cause)
        for r in self.rows:
            if norm_name(r["gbd_cause"]) == key and r["role"] in ("proposed", "registered", "no single descriptor"):
                return r
        return None

    def facts(self, cause: str, expression: str) -> dict | None:
        """The fact row of an expression for a cause (proposal or listed alternative)."""
        key = norm_name(cause)
        for r in self.rows:
            if norm_name(r["gbd_cause"]) == key and r["expression"] == expression:
                return r
        return None


def load_candidates(path: Path | None = None) -> Candidates:
    path = path or CANDIDATES
    rows = read_table(path, "candidates table", CANDIDATE_COLUMNS)
    seen: set[str] = set()
    for i, r in enumerate(rows, 2):
        where = f"{shown(path)}, line {i} ('{r['gbd_cause']}')"
        if r["role"] not in CANDIDATE_ROLES:
            raise GbdError(f"{where}: role must be one of {list(CANDIDATE_ROLES)}, found {r['role']!r}")
        if r["role"] != "alternative":
            if norm_name(r["gbd_cause"]) in seen:
                raise GbdError(f"{where}: a cause can have only one row with the role proposed, registered or no single descriptor")
            seen.add(norm_name(r["gbd_cause"]))
        if r["role"] == "no single descriptor":
            if r["expression"]:
                raise GbdError(f"{where}: a row 'no single descriptor' must not carry an expression")
            continue
        if not EXPRESSION_RE.match(r["expression"]):
            raise GbdError(f"{where}: expression is not of the form \"Heading\"[Mesh] or (\"Heading\"[Mesh] NOT \"Heading\"[Mesh]): "
                           f"{r['expression']!r}")
        if not re.fullmatch(r"D\d+( NOT D\d+)?", r["mesh_ui"]):
            raise GbdError(f"{where}: mesh_ui must be the MeSH unique identifier (D and digits; for the NOT form 'D... NOT D...'), "
                           f"found {r['mesh_ui']!r}")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", r["date_introduced"]):
            raise GbdError(f"{where}: date_introduced must be written YYYY-MM-DD (MeSH Browser, 'Date Introduced', for example "
                           f"1993-01-01), found {r['date_introduced']!r}. If a spreadsheet program changed the dates, restore the file.")
        if r["scope_relation"] not in SCOPE_RELATIONS:
            raise GbdError(f"{where}: scope_relation must be one of {list(SCOPE_RELATIONS)}")
    return Candidates(rows, path)


def proposal_cells(p: dict | None) -> dict:
    """What a rater sheet shows for a cause, from the candidates row (or the lack of one)."""
    if p is None:
        return {"proposed_expression": "", "mesh_ui": "", "date_introduced": "", "scope_relation": "", "mesh_scope_note": "",
                "narrower_descriptors": "", "alternatives": "",
                "inclusion_note": "NO PROPOSAL IN THE CANDIDATES TABLE: look the cause up in the NLM MeSH Browser "
                                  "(meshb.nlm.nih.gov), choose 'use alternative' and write the expression, or choose rule d."}
    return {"proposed_expression": NONE_PROPOSED if p["role"] == "no single descriptor" else p["expression"],
            "mesh_ui": p["mesh_ui"], "date_introduced": p["date_introduced"], "scope_relation": p["scope_relation"],
            "mesh_scope_note": p["mesh_scope_note"], "narrower_descriptors": p["narrower_descriptors"],
            "inclusion_note": p["inclusion_note"], "alternatives": p["alternatives"]}


# ---------------------------------------------------------------------------------------------
# 2. sheets
# ---------------------------------------------------------------------------------------------

def _write_sheet(path: Path, rater: str, intro: list[str], columns: list[str], rows: list[list], link_column: str | None) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping"
    for i, line in enumerate(intro, 1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1))
    ws.cell(row=len(intro) + 1, column=1, value="rater")
    ws.cell(row=len(intro) + 1, column=2, value=rater)
    header_row = len(intro) + 3
    for j, col in enumerate(columns, 1):
        c = ws.cell(row=header_row, column=j, value=col)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF2CC" if col.isupper() else "D9E1F2")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, values in enumerate(rows, 1):
        for j, v in enumerate(values, 1):
            c = ws.cell(row=header_row + i, column=j, value=(v if v != "" else None))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if columns[j - 1].isupper():
                c.fill = PatternFill("solid", fgColor="FFF2CC")
            if link_column and columns[j - 1] == link_column and v:
                c.hyperlink = str(v)
                c.font = Font(color="0563C1", underline="single")
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True, showDropDown=False)
    dv.showErrorMessage = True
    dv.error, dv.errorTitle = "Choose one of the three entries of the list.", "DECISION"
    ws.add_data_validation(dv)
    col_letter = openpyxl.utils.get_column_letter(columns.index("DECISION") + 1)
    dv.add(f"{col_letter}{header_row + 1}:{col_letter}{header_row + max(len(rows), 1)}")
    widths = {"rank": 6, "gbd_cause": 30, "proposed_expression": 40, "mesh_ui": 14, "mesh_browser_link": 44, "date_introduced": 12,
              "scope_relation": 11, "mesh_scope_note": 60, "narrower_descriptors": 60, "inclusion_note": 60, "alternatives": 44,
              "DECISION": 26, "ALTERNATIVE_EXPRESSION": 40, "COMMENT": 40}
    for j, col in enumerate(columns, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = widths.get(col, 36)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def cmd_sheets(args: argparse.Namespace) -> int:
    raters = list(args.raters)
    if len(set(raters)) != 2 or not all(re.fullmatch(r"[A-Za-z0-9]{1,12}", r) for r in raters) or CONSENSUS_KEY in [r.upper() for r in raters]:
        raise GbdError("--raters needs two different keys of letters and digits, for example WP IE")
    ranking_path = Path(args.ranking) if args.ranking else RANKING_PUBLIC
    if not ranking_path.is_file():
        raise GbdError(f"ranking not found: {shown(ranking_path)}. Run 'python scripts/gbd_panel.py rank ...' first.")
    ranking = read_table(ranking_path, "ranking", ["rank", "cause_name", "panel_rule_result"])
    cands = load_candidates(Path(args.candidates) if args.candidates else None)
    kept = sorted((r for r in ranking if r["panel_rule_result"] == G.KEPT), key=lambda r: G.rank_of(r["rank"]))
    if args.n < PANEL_SIZE:
        raise GbdError(f"--n must be at least {PANEL_SIZE}: the panel has {PANEL_SIZE} members and the sheets should hold some reserve")
    selected = kept[: args.n]
    if len(selected) < PANEL_SIZE:
        raise GbdError(f"only {len(selected)} causes of the ranking passed rules (a) to (c); {PANEL_SIZE} are needed")

    out_dir = Path(args.out_dir) if args.out_dir else RATING_DIR
    targets = {rater: out_dir / f"gbd_mapping_rater_{rater}.xlsx" for rater in raters}
    existing = [p for p in targets.values() if p.exists()]
    if existing and not args.force:
        raise GbdError("these sheets exist already and may have been filled in: " + ", ".join(shown(p) for p in existing)
                       + ". Nothing was written. To rate more causes without losing decisions, write new sheets into another "
                         f"folder and carry the decisions over: sheets --n {args.n + 10} --out-dir <new folder> --carry-from {shown(out_dir)}")
    carried: dict[str, dict[str, dict]] = {}
    if args.carry_from:
        for rater in raters:
            old = Path(args.carry_from) / f"gbd_mapping_rater_{rater}.xlsx"
            if old.is_file():
                carried[rater] = read_sheet(old)[1]

    no_proposal = [r["cause_name"] for r in selected if cands.proposal(r["cause_name"]) is None]
    for rater in raters:
        intro = [
            f"GBD panel: mapping of GBD 2023 Level 3 causes to MeSH descriptors. Rater: {rater}",
            "Fill in this sheet alone, without talking to the other rater (Amendment v3.0, Section 5.3). Only the three yellow columns are yours. "
            "Do not change or delete anything else.",
            "DECISION: 'agree' = the proposed expression is the right one. 'use alternative' = write your expression into ALTERNATIVE_EXPRESSION, "
            "in the exact form \"Heading\"[Mesh]. 'no single descriptor (rule d)' = no MeSH heading fits; the cause drops out and the next one moves up.",
            "Rule 5.3: one exploded descriptor per cause whose scope note covers the cause as closely as possible; over- or under-inclusion is stated, "
            "not hidden. A proposal written as (\"A\"[Mesh] NOT \"B\"[Mesh]) removes one descriptor from another; the row explains why.",
            "Check every proposal yourself in the NLM MeSH Browser (column mesh_browser_link): scope note, tree position, narrower descriptors. "
            "The pre-filled facts are a proposal, not a decision.",
            "Stroke, Diabetes mellitus and COVID-19 are named fields of the study: their expression is fixed by rule 5.2, choose 'agree'. "
            "Where no expression is proposed ('none proposed: rule d'), choose 'no single descriptor (rule d)' or 'use alternative'. "
            "Every row needs a DECISION. Save the file under the same name, as .xlsx.",
        ]
        rows = []
        for r in selected:
            cells = proposal_cells(cands.proposal(r["cause_name"]))
            first_ui = re.search(r"D\d+", cells["mesh_ui"] or "")
            old = carried.get(rater, {}).get(norm_name(r["cause_name"]), {})
            rows.append([G.rank_of(r["rank"]), r["cause_name"], cells["proposed_expression"], cells["mesh_ui"],
                         MESH_BROWSER + first_ui.group(0) if first_ui else "", cells["date_introduced"], cells["scope_relation"],
                         cells["mesh_scope_note"], cells["narrower_descriptors"], cells["inclusion_note"], cells["alternatives"],
                         old.get("DECISION", ""), old.get("ALTERNATIVE_EXPRESSION", ""), old.get("COMMENT", "")])
        _write_sheet(targets[rater], rater, intro, SHEET_COLUMNS, rows, "mesh_browser_link")
    print(f"{len(selected)} causes (ranks {selected[0]['rank']} to {selected[-1]['rank']}; {PANEL_SIZE} for the panel, "
          f"{len(selected) - PANEL_SIZE} in reserve) written to:")
    for p in targets.values():
        print(f"  {shown(p)}")
    if carried:
        print(f"decisions carried over from {shown(Path(args.carry_from))} for: {sorted(carried)}")
    if no_proposal:
        print(f"\nNOTE: no proposal in the candidates table for: {no_proposal}. The raters look these up in the MeSH Browser; "
              "a descriptor both choose then needs a row in data/gbd/gbd_mesh_candidates.csv (role alternative; expression, "
              "mesh_ui, date_introduced, scope_relation, inclusion_note).")
    print("\nNext step: each rater fills in his or her own sheet alone; then\n  python scripts/gbd_panel.py reconcile")
    return 0


# ---------------------------------------------------------------------------------------------
# 3. reconcile
# ---------------------------------------------------------------------------------------------

def _cell_text(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):               # a date cell: spreadsheet programs may turn 1993-01-01 into a date
        return v.strftime("%Y-%m-%d")
    return G.one_line(v)


def read_sheet(path: Path) -> tuple[str, dict[str, dict]]:
    """A rater sheet or consensus sheet -> (rater key, {normalised cause name: row}).

    The table is found by its header row (first cell 'rank'), whatever the worksheet is called and
    whatever the order of the rows; columns are read by name."""
    import openpyxl

    if not path.is_file():
        raise GbdError(f"sheet not found: {shown(path)}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise GbdError(f"{shown(path)} cannot be read as an Excel file (.xlsx); save the sheet as .xlsx: {exc}") from exc
    rater, header, out = "", None, {}
    for ws in wb.worksheets:
        for values in ws.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in values]
            if header is None:
                if cells and cells[0] == "rater" and len(cells) > 1:
                    rater = cells[1]
                if cells and cells[0] == "rank" and all(c in cells for c in SHEET_REQUIRED):
                    header = cells
                continue
            if not any(cells):
                continue
            row = dict(zip(header, cells))
            if not row.get("gbd_cause"):
                raise GbdError(f"{path.name}: a row of the table has text but no cause ({[c for c in cells if c][:2]}). No row may be "
                               "emptied or deleted, and remarks belong into the COMMENT column of a cause.")
            key = norm_name(row["gbd_cause"])
            if key in out:
                raise GbdError(f"{path.name}: the cause '{row['gbd_cause']}' occurs twice")
            out[key] = row
        if header is not None:
            break
    if header is None:
        raise GbdError(f"{path.name}: the header row of the mapping sheet was not found; the file was not produced by this script "
                       "or was restructured")
    if not rater:
        raise GbdError(f"{path.name}: the cell next to 'rater' above the table is empty; it must hold the rater key")
    if not out:
        raise GbdError(f"{path.name}: the sheet has no rows")
    return rater, out


def chosen(row: dict, sheet: str) -> str:
    """The decision of one row: an expression, or '' for rule (d). Raises GbdError for a row that
    is empty or contradicts itself; nothing is guessed."""
    d = row["DECISION"].strip().lower()
    proposed = row["proposed_expression"]
    has_proposal = bool(proposed) and proposed != NONE_PROPOSED
    alt = row["ALTERNATIVE_EXPRESSION"]
    for typographic in ("“", "”", "„", "‟", "«", "»"):     # quotation marks that office software may insert
        alt = alt.replace(typographic, '"')
    alt = alt.strip()
    where = f"{sheet}, {row['gbd_cause']}"
    if d == "agree":
        if alt:
            raise GbdError(f"{where}: DECISION is 'agree' but ALTERNATIVE_EXPRESSION is filled in. Choose 'use alternative' "
                           "or empty the alternative.")
        if not has_proposal:
            raise GbdError(f"{where}: there is no proposed expression to agree with. Choose 'use alternative' and write an "
                           "expression, or 'no single descriptor (rule d)'.")
        return proposed
    if d == "use alternative":
        if not EXPRESSION_RE.match(alt):
            raise GbdError(f"{where}: ALTERNATIVE_EXPRESSION {alt!r} is not of the form "
                           "\"Heading\"[Mesh] or (\"Heading\"[Mesh] NOT \"Heading\"[Mesh])")
        if alt == proposed:
            raise GbdError(f"{where}: the alternative equals the proposal; choose 'agree'")
        return alt
    if d.startswith("no single descriptor"):
        if alt:
            raise GbdError(f"{where}: DECISION is 'no single descriptor' but ALTERNATIVE_EXPRESSION is filled in")
        return ""
    raise GbdError(f"{where}: DECISION is empty or not one of {list(DECISIONS)}")


def _collect(problems: list[str], func, *args):
    try:
        return func(*args)
    except (GbdError, G.GbdPanelError) as exc:
        problems.append(str(exc))
        return None


def write_consensus_sheet(path: Path, r1: str, r2: str, rows: list[dict]) -> None:
    intro = [
        "GBD panel: consensus sheet for the rows on which the two raters differ (Amendment v3.0, Section 5.3: resolved by discussion).",
        f"Discuss each row together ({r1} and {r2}) and enter the agreed decision in the yellow columns: 'agree' = the proposed expression; "
        "'use alternative' = write the agreed expression into ALTERNATIVE_EXPRESSION; 'no single descriptor (rule d)' = the cause drops out.",
        "Write the reason for the agreed decision into COMMENT. Every row needs a DECISION. Save under the same name, then run reconcile again.",
    ]
    columns = ["rank", "gbd_cause", "proposed_expression", f"choice_{r1}", f"comment_{r1}", f"choice_{r2}", f"comment_{r2}",
               "alternatives", "DECISION", "ALTERNATIVE_EXPRESSION", "COMMENT"]
    values = [[G.rank_of(r["rank"]), r["gbd_cause"], r["proposed_expression"], r[f"choice_{r1}"], r[f"comment_{r1}"], r[f"choice_{r2}"],
               r[f"comment_{r2}"], r["alternatives"], "", "", ""] for r in rows]
    _write_sheet(path, CONSENSUS_KEY, intro, columns, values, None)


def cmd_reconcile(args: argparse.Namespace) -> int:
    if args.sheets:
        paths = [Path(p) for p in args.sheets]
    else:
        paths = sorted(RATING_DIR.glob("gbd_mapping_rater_*.xlsx")) if RATING_DIR.is_dir() else []
        if len(paths) != 2:
            raise GbdError(f"expected exactly two rater sheets in {shown(RATING_DIR)} (gbd_mapping_rater_<key>.xlsx), found "
                           f"{[p.name for p in paths]}. Name the two files with --sheets.")
    # the two raters in the alphabetical order of their keys, whatever the order of the files
    (r1, s1), (r2, s2) = sorted((read_sheet(paths[0]), read_sheet(paths[1])), key=lambda pair: pair[0])
    if CONSENSUS_KEY in (r1.upper(), r2.upper()):
        raise GbdError("one of the two files is the consensus sheet; --sheets takes the two rater sheets")
    if r1 == r2:
        raise GbdError(f"both sheets name the same rater ({r1}); two different raters are required")
    if set(s1) != set(s2):
        only1 = sorted(s1[k]["gbd_cause"] for k in set(s1) - set(s2))
        only2 = sorted(s2[k]["gbd_cause"] for k in set(s2) - set(s1))
        raise GbdError(f"the two sheets do not list the same causes (only in the sheet of {r1}: {only1}; only in the sheet of {r2}: "
                       f"{only2}). Both must come from the same 'sheets' run and no row may be deleted or added.")

    # 1. The sheets against the ranking: the rated causes are the first causes that passed rules (a) to (c),
    #    with the ranks of the ranking. A sheet in which ranks were changed, or a ranking that was made
    #    again from another export, is refused.
    ranking_path = Path(args.ranking) if args.ranking else RANKING_PUBLIC
    if not ranking_path.is_file():
        raise GbdError(f"ranking not found: {shown(ranking_path)}. Run 'python scripts/gbd_panel.py rank ...' first.")
    ranking = read_table(ranking_path, "ranking", ["rank", "cause_id", "cause_name", "level1_group", "panel_rule_result"])
    problems: list[str] = []
    for rater, sheet in ((r1, s1), (r2, s2)):
        problems += [f"sheet of {rater}: {p}" for p in G.problems_against_ranking(list(sheet.values()), ranking)]
    if problems:
        raise GbdError("the rater sheets do not match the ranking " + shown(ranking_path) + ":\n  - " + "\n  - ".join(problems[:12])
                       + "\nIf the ranking was made again from another export, new sheets are needed (sheets --out-dir <new folder> "
                         "--carry-from <old folder>). If a sheet was changed outside the yellow columns, restore it.")

    # 2. The sheets against the candidates table: what the raters saw is what the table says now.
    cands = load_candidates(Path(args.candidates) if args.candidates else None)
    for rater, sheet in ((r1, s1), (r2, s2)):
        for row in sheet.values():
            expected = proposal_cells(cands.proposal(row["gbd_cause"]))
            columns = ["proposed_expression"] + (FACT_COLUMNS if expected["proposed_expression"] else [])
            for col in columns:
                if col in row and G.one_line(row[col]) != G.one_line(expected[col]):
                    problems.append(f"sheet of {rater}, {row['gbd_cause']}: column {col} differs from the candidates table")
    if problems:
        raise GbdError("the pre-filled part of a sheet differs from " + shown(cands.path) + ":\n  - " + "\n  - ".join(problems[:12])
                       + "\nEither a sheet was changed outside the yellow columns, or the candidates table was changed after the "
                         "sheets were written. Restore the one that changed.")

    # 3. Every decision, of both raters, before anything else is said.
    order = sorted(s1, key=lambda k: G.rank_of(s1[k]["rank"]))
    choice: dict[str, tuple[str | None, str | None]] = {}
    for key in order:
        choice[key] = (_collect(problems, chosen, s1[key], f"sheet of {r1}"), _collect(problems, chosen, s2[key], f"sheet of {r2}"))
    if problems:
        raise GbdError(f"{len(problems)} row(s) cannot be read as a decision:\n  - " + "\n  - ".join(problems))

    # 4. Named fields keep their expression (Section 5.2).
    for key in order:
        p = cands.proposal(s1[key]["gbd_cause"])
        if p is not None and p["role"] == "registered":
            for rater, c in zip((r1, r2), choice[key]):
                if c != p["expression"]:
                    problems.append(f"sheet of {rater}, {s1[key]['gbd_cause']}: this cause is a named field of the study; its expression "
                                    f"{p['expression']} is fixed by Section 5.2 of the amendment. Choose 'agree'.")
    if problems:
        raise GbdError("named fields:\n  - " + "\n  - ".join(problems))

    # 5. Disagreements and the consensus sheet.
    differing = [k for k in order if choice[k][0] != choice[k][1]]
    consensus_path = Path(args.consensus) if args.consensus else paths[0].parent / "gbd_mapping_consensus.xlsx"
    consensus: dict[str, str] = {}
    consensus_comment: dict[str, str] = {}
    label = lambda c: c or "rule (d)"   # noqa: E731
    if differing:
        template_rows = [{"rank": s1[k]["rank"], "gbd_cause": s1[k]["gbd_cause"], "proposed_expression": s1[k]["proposed_expression"],
                          f"choice_{r1}": label(choice[k][0]), f"comment_{r1}": s1[k]["COMMENT"], f"choice_{r2}": label(choice[k][1]),
                          f"comment_{r2}": s2[k]["COMMENT"], "alternatives": s1[k].get("alternatives", "")} for k in differing]
        usable = False
        if consensus_path.is_file():
            rater_c, sc = read_sheet(consensus_path)
            if rater_c.upper() != CONSENSUS_KEY:
                raise GbdError(f"{shown(consensus_path)} is not a consensus sheet (its rater cell says {rater_c!r}). A rater sheet "
                               "cannot serve as the consensus; fill in the consensus sheet that reconcile writes.")
            stale = set(sc) != set(differing) or any(
                sc[k].get(f"choice_{r1}") != label(choice[k][0]) or sc[k].get(f"choice_{r2}") != label(choice[k][1]) for k in differing)
            if stale:
                raise GbdError(f"{shown(consensus_path)} does not show the rows and choices on which the raters differ now "
                               f"({[s1[k]['gbd_cause'] for k in differing]}). The rater sheets changed since it was written: "
                               "delete the consensus sheet and run reconcile again.")
            if all(sc[k]["DECISION"].strip() for k in differing):
                for k in differing:
                    c = _collect(problems, chosen, sc[k], "consensus sheet")
                    consensus[k] = c if c is not None else ""
                    consensus_comment[k] = sc[k]["COMMENT"]
                if problems:
                    raise GbdError("the consensus sheet cannot be read:\n  - " + "\n  - ".join(problems))
                usable = True
        else:
            write_consensus_sheet(consensus_path, r1, r2, template_rows)
        if not usable:
            print(f"{len(differing)} disagreement(s) between {r1} and {r2}. No panel was written.\n")
            for k in differing:
                a, b = s1[k], s2[k]
                print(f"  rank {a['rank']}: {a['gbd_cause']}\n      {r1}: {label(choice[k][0])}{'   (' + a['COMMENT'] + ')' if a['COMMENT'] else ''}"
                      f"\n      {r2}: {label(choice[k][1])}{'   (' + b['COMMENT'] + ')' if b['COMMENT'] else ''}")
            print(f"\nResolve by discussion (Section 5.3): open {shown(consensus_path)}, enter the agreed DECISION for each of these "
                  "rows, save, and run the same reconcile command again.")
            return 1
        print(f"{len(differing)} disagreement(s), resolved by the consensus sheet {shown(consensus_path)}:")
        for k in differing:
            print(f"  rank {s1[k]['rank']}: {s1[k]['gbd_cause']}: {label(consensus[k])}")
        print()
    elif args.consensus:
        print(f"NOTE: the raters agree on every row; the consensus sheet {shown(consensus_path)} was not needed and not read.\n")

    # 6. The mapping table.
    by_name = {norm_name(r["cause_name"]): r for r in ranking}
    rows = []
    for key in order:
        a, b = s1[key], s2[key]
        ca, cb = choice[key]
        final, how = (ca, G.RECONCILIATION_VALUES[0]) if ca == cb else (consensus[key], G.RECONCILIATION_VALUES[1])
        rk = by_name[key]
        row = {"rank": G.rank_of(rk["rank"]), "cause_id": rk["cause_id"], "gbd_cause": rk["cause_name"], "level1_group": rk["level1_group"],
               f"choice_{r1}": label(ca), f"choice_{r2}": label(cb), "reconciliation": how,
               f"comment_{r1}": a["COMMENT"], f"comment_{r2}": b["COMMENT"], "comment_consensus": consensus_comment.get(key, "")}
        if final:
            facts = cands.facts(rk["cause_name"], final)
            if facts is None:
                raise GbdError(f"'{rk['cause_name']}': the agreed expression {final} has no row for this cause in {shown(cands.path)}. "
                               "Add one line with the role 'alternative': gbd_cause, expression, mesh_ui (MeSH Browser, 'Unique ID'), "
                               "date_introduced (MeSH Browser, 'Date Introduced', written YYYY-MM-DD), scope_relation (exact, broader, "
                               "narrower or mixed) and inclusion_note. Edit the file with a text editor, not with a spreadsheet "
                               "program, and run reconcile again. Nothing was written.")
            row.update({"decision": G.MAPPED, "expression": final, "mesh_ui": facts["mesh_ui"], "date_introduced": facts["date_introduced"],
                        "scope_relation": facts["scope_relation"], "inclusion_note": G.one_line(facts["inclusion_note"])})
        else:
            row.update({"decision": G.RULE_D, "expression": "", "mesh_ui": "", "date_introduced": "", "scope_relation": "",
                        "inclusion_note": "no single MeSH descriptor with a corresponding scope note"})
        rows.append(row)
    columns = ["rank", "cause_id", "gbd_cause", "level1_group", "decision", "expression", "mesh_ui", "date_introduced", "scope_relation",
               "inclusion_note", f"choice_{r1}", f"choice_{r2}", "reconciliation", f"comment_{r1}", f"comment_{r2}", "comment_consensus"]
    mapping_path = Path(args.mapping) if args.mapping else MAPPING
    text = csv_text(columns, [{k: G.one_line(v) for k, v in r.items()} for r in rows])
    n_mapped = sum(1 for r in rows if r["decision"] == G.MAPPED)
    if n_mapped < PANEL_SIZE:
        raise GbdError(f"only {n_mapped} of the {len(rows)} rated causes are mapped; the panel needs {PANEL_SIZE}. More causes have to be "
                       f"rated. Write larger sheets that keep the decisions already made:\n  python scripts/gbd_panel.py sheets "
                       f"--n {len(rows) + 10} --out-dir {shown(paths[0].parent)}_2 --carry-from {shown(paths[0].parent)}\n"
                       "The raters fill in the new rows; then run reconcile with --sheets on the two new files. Nothing was written.")
    n_independent = sum(1 for r in rows if r["reconciliation"] == G.RECONCILIATION_VALUES[0])
    print(f"{len(rows)} causes reconciled: {n_independent} agreed independently, {len(rows) - n_independent} after discussion.")
    return build_panel(G.read_csv_dicts_from_text(text), "final", mapping_text=text, mapping_path=mapping_path, candidates_path=cands.path)


# ---------------------------------------------------------------------------------------------
# 4. build
# ---------------------------------------------------------------------------------------------

def yaml_quote(s: str) -> str:
    return "'" + G.one_line(s).replace("'", "''") + "'"


def build_panel(rows: list[dict], status: str, *, mapping_text: str, mapping_path: Path, candidates_path: Path | None = None,
                panel_yaml: Path | None = None, table_md: Path | None = None, provenance: str = "", test_fixture: bool = False) -> int:
    """Write the panel file and the table for the amendment from the rows of a mapping table.

    Everything is checked before anything is written: the panel file is written to a temporary
    name, read back with the pipeline's own loader, and only then moved into place together with
    the mapping table. A final panel needs a reconciled mapping table, the ranking and the export
    record, and must agree with all three."""
    import yaml

    from airct_benchmark.catalogue import CatalogueError, load_catalogue
    from airct_benchmark.queries import build_amendment_count_queries, build_count_queries, shared_cells

    panel_yaml = panel_yaml or PANEL_YAML
    table_md = table_md or AMENDMENT_TABLE
    root = panel_yaml.resolve().parent.parent
    registered = yaml.safe_load(CATALOGUE_YAML.read_text(encoding="utf-8"))
    fields = registered["fields"]
    field_by_mesh = {f["mesh"]: k for k, f in fields.items()}
    windows = {k: v for k, v in registered["years"].items() if k != "default"}      # named fields with a shorter year window
    final = status == "final"
    strict = final and not test_fixture

    raters: tuple[str, str] | None = None
    if strict:
        raters = G.check_reconciled(rows)
    elif final:
        raters = ("WP", "IE")           # synthetic panel of the test suite
    members = G.members_from_mapping(rows, field_by_mesh)
    if final and len(members) < PANEL_SIZE:
        raise GbdError(f"only {len(members)} causes are mapped; a final panel needs {PANEL_SIZE}")
    if not members:
        raise GbdError("the mapping table has no mapped cause")
    last_rank = members[-1]["rank"]

    ranking = read_table(RANKING_PUBLIC, "ranking", ["rank", "cause_name", "panel_rule_result"]) if RANKING_PUBLIC.is_file() else []
    record = json.loads(EXPORT_RECORD.read_text(encoding="utf-8")) if EXPORT_RECORD.is_file() else {}
    if strict:
        if not ranking or not record:
            raise GbdError(f"a final panel needs the ranking ({shown(RANKING_PUBLIC)}) and the export record ({shown(EXPORT_RECORD)}); "
                           "run 'rank' first")
        problems = G.problems_against_ranking(rows, ranking)
        if (record.get("ranking_public") or {}).get("sha256") != sha256_text(RANKING_PUBLIC):
            problems.append("the export record names another ranking than the one in data/gbd; run 'rank' again")
        if problems:
            raise GbdError("the mapping table does not match the ranking:\n  - " + "\n  - ".join(problems[:12]))
    ranking_sha = sha256_text(RANKING_PUBLIC) if ranking else ""
    mapping_sha = sha256_of_text(mapping_text)

    lines = [
        "# GBD panel of the AI-RCT Benchmark (Protocol Amendment v3.0, Section 5, decision D24).",
        "# GENERATED by scripts/gbd_panel.py from the mapping table. Do not edit by hand: change the",
        "# rater sheets and reconcile again, so that this file, the table of the amendment and the hashes agree.",
        "#",
        "# status 'final'       : export, two independent rater sheets and reconciliation are complete;",
        "#                        the pipeline runs with this panel.",
        "# status 'provisional' : placeholder for development and code review; freeze, lists and",
        "#                        syntax-check refuse to run.",
        f"status: {status}",
    ]
    if test_fixture:
        # synthetic panel of the test suite: the online modes of the pipeline refuse to send anything with it
        lines.append("test_fixture: true")
    if provenance:
        lines.append(f"provenance: {yaml_quote(provenance)}")
    lines += [
        "source:",
        f"  gbd_round: {yaml_quote('GBD 2023')}",
        f"  citation: {yaml_quote(GBD_CITATION)}",
        f"  quantity: {yaml_quote('Level 3 causes, DALYs (number), all ages, both sexes, Global, 2023')}",
        f"  selection_rule: {yaml_quote('Amendment v3.0, Section 5.2: rank by DALY count; exclude (a) Injuries, (b) pandemic-related outcome categories, (c) residual causes whose name begins with Other, (d) causes without a single MeSH descriptor; first 20 remaining causes')}",
        f"  mapping_table: {yaml_quote(rel(mapping_path, root))}",
        f"  mapping_table_sha256: {yaml_quote(mapping_sha)}",
        f"  ranking_sha256: {yaml_quote(ranking_sha)}",
        f"  raw_export_sha256: {yaml_quote((record.get('raw_export') or {}).get('sha256', ''))}",
        f"  candidates_sha256: {yaml_quote(sha256_text(candidates_path) if candidates_path and candidates_path.is_file() else '')}",
        f"  raters: {yaml_quote(' and '.join(raters) if raters else '')}",
        "members:",
    ]
    for m in members:
        lines += [
            f"  - rank: {m['rank']}",
            f"    key: {m['key']}",
            f"    cause: {yaml_quote(m['cause'])}",
            f"    expression: {yaml_quote(m['expression'])}",
            f"    mesh_ui: {yaml_quote(m['mesh_ui'])}",
            f"    date_introduced: {yaml_quote(m['date_introduced'])}",
            f"    in_trajectory: {'true' if m['in_trajectory'] else 'false'}",
            f"    shared_with: {m['shared_with'] if m['shared_with'] else 'null'}",
            f"    scope_relation: {yaml_quote(m['scope_relation'])}",
            f"    inclusion_note: {yaml_quote(m['inclusion_note'])}",
        ]
    panel_text = "\n".join(lines) + "\n"

    # the pipeline's own loader is the judge, before the file takes its place
    tmp = panel_yaml.with_name(panel_yaml.name + ".tmp")
    write_text_lf(tmp, panel_text)
    try:
        cat = load_catalogue(CATALOGUE_YAML, tmp)
        shared_cells(cat)
        n_new, n_all = len(build_amendment_count_queries(cat)), len(build_count_queries(cat))
        panel_sha = cat.gbd_sha256()
    except CatalogueError as exc:
        tmp.unlink(missing_ok=True)
        raise GbdError(f"the pipeline does not accept this panel, nothing was written: {exc}") from exc
    try:
        if not mapping_path.is_file() or sha256_text(mapping_path) != mapping_sha:
            write_text_lf(mapping_path, mapping_text)
        os.replace(tmp, panel_yaml)
    finally:
        tmp.unlink(missing_ok=True)      # nothing is left behind when a file could not be written (for example: open in Excel)

    # table for the amendment (Section 5): one row per member, exactly as the identity check expects it
    left_out = []                  # (rank, text), listed in rank order whatever the rule
    for r in ranking:
        if G.rank_of(r["rank"]) <= last_rank and r["panel_rule_result"] != G.KEPT:
            left_out.append((G.rank_of(r["rank"]), f"{r['cause_name']} (rank {r['rank']}, {r.get('exclusion_reason', '')})"))
    for r in rows:
        if r["decision"] != G.MAPPED and G.rank_of(r["rank"]) <= last_rank:
            left_out.append((G.rank_of(r["rank"]), f"{r['gbd_cause']} (rank {r['rank']}, rule (d): no single MeSH descriptor)"))
    excluded = [text for _, text in sorted(left_out)]
    if not ranking:
        excluded_text = "not stated (no ranking file)."
    else:
        excluded_text = "; ".join(excluded) + "." if excluded else "none."
    md = [
        f"<!-- generated by scripts/gbd_panel.py; panel status: {status}; paste into Amendment v3.0, Section 5, unchanged -->",
        "",
        "5.7 GBD panel" + (" (final)." if final else " (PROVISIONAL, not for the OSF update)."),
        f"The {len(members)} GBD 2023 Level 3 causes of the panel in rank order of global DALY counts 2023, selected by the rule of 5.2 and "
        "mapped by the rule of 5.3" + (f" by {raters[0]} and {raters[1]}" if raters else "") + ". The query expression is used verbatim as "
        "{G} in families GA, GD and GE. Members that coincide with a named field share its registered cells and are not queried again. "
        + "".join(f"{m['cause']} coincides with the named field {m['shared_with']}, whose registered year window is "
                  f"{windows[m['shared_with']]['start']} to {windows[m['shared_with']]['end']}: it shares the cells of that window and has "
                  "no cell for earlier years. " for m in members if m["shared_with"] in windows)
        + "Descriptors introduced after 2015 enter 2024 and the pooled windows only (panel rule of Q13). Scope relation: exact, broader or "
        "narrower than the GBD cause, or mixed where over- and under-inclusion are both stated.",
        "",
        G.TABLE_HEADER,
        G.TABLE_RULE,
    ] + [G.table_row(m) for m in members] + [
        "",
        "Causes ranked above the last panel member that are not in the panel: " + excluded_text,
        "",
    ]
    if record:
        md += [f"Export: GBD Results Tool, {record.get('export_date', '')}, by {record.get('exported_by', '')}, permalink {record.get('permalink', '')}; "
               f"SHA-256 of the raw export {(record.get('raw_export') or {}).get('sha256', '')}, of the cause hierarchy file "
               f"{(record.get('cause_hierarchy') or {}).get('sha256', '')}, of the deposited ranking {ranking_sha}, of the mapping table "
               f"{mapping_sha}, of config/gbd_panel.yaml {panel_sha} (deposited text files: hash of the content with LF "
               f"line endings). The raw export and the hierarchy file are kept by the authors and are not redistributed. {IHME_SOURCE_LINE}", ""]
    md += [G.volume_line(n_new, n_all), ""]
    write_text_lf(table_md, "\n".join(md))

    if strict:
        problems = G.panel_file_problems(yaml.safe_load(panel_text), root, field_by_mesh)
        if problems:
            raise GbdError("the panel was written, but it does not agree with the deposited files:\n  - " + "\n  - ".join(problems))

    shared = [m["cause"] for m in members if m["shared_with"]]
    late = [m["cause"] for m in members if not m["in_trajectory"]]
    print(f"panel written: {shown(panel_yaml)} (status {status}, {len(members)} members, ranks {members[0]['rank']} to {last_rank})")
    print(f"mapping table: {shown(mapping_path)}")
    print(f"  shared with a named field, not queried again: {shared or 'none'}")
    print(f"  descriptor introduced after 2015 (2024 and pooled windows only): {late or 'none'}")
    print(f"  catalogue volume: {n_new:,} additional count requests, {n_all:,} per run")
    print(f"table for the amendment written: {shown(table_md)}")
    if final:
        print("\nNext steps:\n"
              "  1. paste the table into the amendment (Section 5), unchanged, and copy the amendment text to docs/AMENDMENT_v3_0.md\n"
              "  2. python -m airct_benchmark check-catalogue      (must end with exit code 0)\n"
              "  3. python -m pytest")
    return 0


def cmd_build(args: argparse.Namespace) -> int:
    mapping_path = Path(args.mapping) if args.mapping else MAPPING
    rows = read_table(mapping_path, "mapping table", ["rank", "gbd_cause", "decision", "expression"])
    text = mapping_path.read_text(encoding="utf-8")
    return build_panel(rows, args.status, mapping_text=text, mapping_path=mapping_path, candidates_path=CANDIDATES,
                       provenance=args.provenance or "")


# ---------------------------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------------------------

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    r = sub.add_parser("rank", help="rank the Level 3 causes of the export and apply rules (a) to (c)")
    r.add_argument("--export", required=True, help="CSV downloaded from the GBD Results Tool (unpacked, untouched)")
    r.add_argument("--hierarchy", required=True, help="GBD 2023 cause hierarchy file from GHDx (xlsx or csv)")
    r.add_argument("--author", required=True, help="who made the export")
    r.add_argument("--export-date", required=True, help="date of the export, YYYY-MM-DD")
    r.add_argument("--permalink", required=True, help="permalink of the Results Tool query")
    r.add_argument("--citation-file", default=None, help="the citation text file that comes with the download (recorded verbatim)")
    r.add_argument("--local-dir", default=None, help="folder for the ranking with DALY values (default: the folder of the raw export)")
    r.add_argument("--show", type=int, default=30, help="how many kept causes to print")
    r.set_defaults(func=cmd_rank)

    s = sub.add_parser("sheets", help="write one pre-filled mapping sheet per rater")
    s.add_argument("--ranking", default=None, help="default: data/gbd/gbd2023_level3_ranking.csv")
    s.add_argument("--candidates", default=None, help="default: data/gbd/gbd_mesh_candidates.csv")
    s.add_argument("--raters", nargs=2, default=["WP", "IE"], help="rater keys (default: WP IE)")
    s.add_argument("--n", type=int, default=DEFAULT_ROWS, help=f"number of causes to rate (default {DEFAULT_ROWS}: 20 for the panel plus reserve)")
    s.add_argument("--out-dir", default=None, help="default: data/gbd/rating")
    s.add_argument("--carry-from", default=None, help="folder with earlier sheets whose decisions are copied into the new sheets")
    s.add_argument("--force", action="store_true", help="overwrite existing sheets (filled sheets are lost)")
    s.set_defaults(func=cmd_sheets)

    c = sub.add_parser("reconcile", help="compare the two filled sheets; write mapping table and final panel when the raters agree")
    c.add_argument("--sheets", nargs=2, default=None, help="the two rater sheets (default: the two files in data/gbd/rating)")
    c.add_argument("--consensus", default=None, help="consensus sheet (default: gbd_mapping_consensus.xlsx next to the rater sheets)")
    c.add_argument("--candidates", default=None, help="default: data/gbd/gbd_mesh_candidates.csv")
    c.add_argument("--ranking", default=None, help="default: data/gbd/gbd2023_level3_ranking.csv")
    c.add_argument("--mapping", default=None, help="default: data/gbd/gbd_panel_mesh_mapping.csv")
    c.set_defaults(func=cmd_reconcile)

    b = sub.add_parser("build", help="rebuild config/gbd_panel.yaml and the amendment table from a mapping table")
    b.add_argument("--mapping", default=None, help="default: data/gbd/gbd_panel_mesh_mapping.csv")
    b.add_argument("--status", choices=["final", "provisional"], default="provisional",
                   help="final needs the reconciled mapping table written by 'reconcile' (default: provisional)")
    b.add_argument("--provenance", default=None, help="free text stored in the panel file (used for provisional panels)")
    b.set_defaults(func=cmd_build)
    return p


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):      # a console that cannot show a character must not stop the run
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    args = build_parser().parse_args(argv)
    try:
        return args.func(args)
    except (GbdError, G.GbdPanelError) as exc:
        print(f"\nSTOPPED: {exc}", file=sys.stderr)
    except PermissionError as exc:
        print(f"\nSTOPPED: a file cannot be read or written: {exc.filename}. It is probably open in Excel or another program; "
              "close it there and run the command again.", file=sys.stderr)
    except FileNotFoundError as exc:
        print(f"\nSTOPPED: file or folder not found: {exc.filename}", file=sys.stderr)
    except (OSError, UnicodeDecodeError, zipfile.BadZipFile) as exc:
        print(f"\nSTOPPED: a file could not be read ({type(exc).__name__}: {exc}). Check that the right file was named and that it "
              "is the untouched download or a sheet saved as .xlsx.", file=sys.stderr)
    except KeyError as exc:
        print(f"\nSTOPPED: a file lacks the column {exc}; it is not the file this step expects.", file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.exit(main())
