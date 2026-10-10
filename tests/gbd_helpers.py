"""Synthetic GBD material for the test suite: an export shaped like the CSV of the GBD Results
Tool, a cause hierarchy, filled rater sheets, and a complete run of scripts/gbd_panel.py into a
temporary directory. The DALY values are invented (descending round numbers); they are not GBD
estimates. Nothing here touches the files of the study or the network."""

from __future__ import annotations

import csv
import importlib.util
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
FIXTURES = Path(__file__).parent / "fixtures"

# Level 3 causes of the synthetic export in descending order of their invented value.
LEVEL3 = [
    ("Ischemic heart disease", "NCD"), ("Neonatal disorders", "CMNN"), ("Stroke", "NCD"),
    ("Lower respiratory infections", "CMNN"), ("Diabetes mellitus", "NCD"), ("Road injuries", "INJ"),
    ("Chronic obstructive pulmonary disease", "NCD"), ("Falls", "INJ"), ("Low back pain", "NCD"),
    ("Diarrheal diseases", "CMNN"), ("Depressive disorders", "NCD"), ("Anxiety disorders", "NCD"),
    ("Congenital birth defects", "NCD"), ("Age-related and other hearing loss", "NCD"), ("Malaria", "CMNN"),
    ("Tracheal, bronchus, and lung cancer", "NCD"), ("Chronic kidney disease", "NCD"), ("Headache disorders", "NCD"),
    ("Other musculoskeletal disorders", "NCD"), ("HIV/AIDS", "CMNN"), ("Tuberculosis", "CMNN"),
    ("Cirrhosis and other chronic liver diseases", "NCD"), ("Alzheimer's disease and other dementias", "NCD"),
    ("Dietary iron deficiency", "CMNN"), ("Self-harm", "INJ"), ("Other COVID-19 pandemic-related outcomes", "CMNN"),
    ("Hypertensive heart disease", "NCD"), ("Colon and rectum cancer", "NCD"), ("Stomach cancer", "NCD"),
    ("Synthetic cause without proposal", "NCD"), ("Breast cancer", "NCD"), ("Asthma", "NCD"), ("Neck pain", "NCD"),
    ("Drug use disorders", "NCD"), ("Osteoarthritis", "NCD"), ("Meningitis", "CMNN"),
]
LEVEL1 = {"CMNN": ("295", "Communicable, maternal, neonatal, and nutritional diseases"),
          "NCD": ("409", "Non-communicable diseases"), "INJ": ("687", "Injuries")}
LEVEL2 = {"CMNN": ("956", "Synthetic level 2 group A"), "NCD": ("491", "Synthetic level 2 group B"), "INJ": ("688", "Synthetic level 2 group C")}
RULE_D = "no single descriptor (rule d)"
EXPORT_NAME = "IHME-GBD_2023_DATA-synthetic.csv"
HIERARCHY_NAME = "IHME_GBD_2023_CAUSE_HIERARCHY_synthetic.xlsx"


def level3_value(index: int) -> int:
    return (len(LEVEL3) - index) * 1000


def level3_total(values: dict[str, float] | None = None) -> float:
    """Sum of the invented Level 3 values, with some of them replaced: the value 'All causes' must have."""
    return sum((values or {}).get(name, level3_value(i)) for i, (name, _) in enumerate(LEVEL3))


def load_script(tmp_path: Path | None = None):
    """A fresh instance of scripts/gbd_panel.py. With ``tmp_path`` every output path is redirected
    into it, laid out like a repository (config/, data/gbd/)."""
    spec = importlib.util.spec_from_file_location("gbd_panel_under_test", ROOT / "scripts" / "gbd_panel.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    if tmp_path is not None:
        base = tmp_path / "data" / "gbd"
        mod.GBD_DIR = base
        mod.RANKING_PUBLIC = base / "gbd2023_level3_ranking.csv"
        mod.EXPORT_RECORD = base / "gbd_export_record.json"
        mod.MAPPING = base / "gbd_panel_mesh_mapping.csv"
        mod.AMENDMENT_TABLE = base / "gbd_panel_amendment_table.md"
        mod.RATING_DIR = base / "rating"
        mod.PANEL_YAML = tmp_path / "config" / "gbd_panel.yaml"
    return mod


def hierarchy_rows() -> list[dict]:
    rows = [{"Cause ID": "294", "Cause Name": "All causes", "Parent ID": "294", "Level": "0"}]
    for cid, name in LEVEL1.values():
        rows.append({"Cause ID": cid, "Cause Name": name, "Parent ID": "294", "Level": "1"})
    for group, (cid, name) in LEVEL2.items():
        rows.append({"Cause ID": cid, "Cause Name": name, "Parent ID": LEVEL1[group][0], "Level": "2"})
    for i, (name, group) in enumerate(LEVEL3):
        rows.append({"Cause ID": str(2000 + i), "Cause Name": name, "Parent ID": LEVEL2[group][0], "Level": "3"})
    rows.append({"Cause ID": "3001", "Cause Name": "Ischemic stroke", "Parent ID": "2002", "Level": "4"})
    return rows


def write_hierarchy(path: Path, *, columns: tuple[str, ...] = ("Cause ID", "Cause Name", "Parent ID", "Level"), rename: dict[str, str] | None = None,
                    extra: list[dict] | None = None, drop: tuple[str, ...] = ()) -> Path:
    rows = [{**r, "Cause Name": (rename or {}).get(r["Cause Name"], r["Cause Name"])} for r in hierarchy_rows() + (extra or [])
            if r["Cause Name"] not in drop]
    if path.suffix.lower() == ".xlsx":
        wb = openpyxl.Workbook()
        wb.active.title = "Notes"
        wb.active.append(["Synthetic hierarchy of the test suite"])
        ws = wb.create_sheet("Cause Hierarchy")
        ws.append(["A title row above the table"])
        ws.append(["Cause ID", "Cause Name", "Parent ID", "Parent Name", "Level", "Cause Outline", "Sort Order"])
        for r in rows:
            ws.append([float(r["Cause ID"]), r["Cause Name"], float(r["Parent ID"]), "", float(r["Level"]), "", ""])
        wb.save(path)
    else:
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(columns), extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
    return path


def export_rows(*, drop: tuple[str, ...] = (), spelling: dict[str, str] | None = None, values: dict[str, float] | None = None) -> list[dict]:
    """Rows shaped like the Results Tool CSV, including rows of other strata that the filter must ignore."""
    out = []
    total = sum(level3_value(i) for i in range(len(LEVEL3)))
    for r in hierarchy_rows():
        name = r["Cause Name"]
        if name in drop:
            continue
        index = next((i for i, (n, _) in enumerate(LEVEL3) if n == name), None)
        val = total if r["Level"] == "0" else (12345 if index is None else level3_value(index))
        val = (values or {}).get(name, val)
        number = val if isinstance(val, (int, float)) else 0
        base = {"measure_id": 2, "measure_name": "DALYs (Disability-Adjusted Life Years)", "location_id": 1, "location_name": "Global",
                "sex_id": 3, "sex_name": "Both", "age_id": 22, "age_name": "All ages", "cause_id": r["Cause ID"],
                "cause_name": (spelling or {}).get(name, name), "metric_id": 1, "metric_name": "Number", "year": 2023,
                "val": val, "upper": number * 1.1, "lower": number * 0.9}
        out.append(base)
        # rows of other strata, with values that would reverse the ranking if they were not filtered out
        other = 1_000_000 - number
        out.append({**base, "metric_id": 3, "metric_name": "Rate", "val": other})
        out.append({**base, "sex_id": 1, "sex_name": "Male", "val": other})
        out.append({**base, "age_id": 27, "age_name": "Age-standardized", "val": other})
        out.append({**base, "year": 2021, "val": other})
        out.append({**base, "measure_id": 1, "measure_name": "Deaths", "val": other})
    return out


GERMAN = {"Both": "Beide", "Male": "Männlich", "All ages": "Alle Altersgruppen", "Age-standardized": "Altersstandardisiert",
          "Number": "Anzahl", "Rate": "Rate", "Deaths": "Todesfälle", "All causes": "Alle Ursachen"}


def write_export(path: Path, *, with_ids: bool = True, delimiter: str = ",", extra_rows: list[dict] | None = None,
                 german: bool = False, population_groups: tuple[int, ...] = (), **kw) -> Path:
    """``german``: names as the German interface of the Results Tool writes them (the ids stay), cause names replaced by
    German-looking ones. ``population_groups``: the column pair of the GBD 2023 tool, one copy of every row per group;
    group 1 is the whole population, any other group carries values that would change the ranking."""
    cols = ["measure_id", "measure_name", "location_id", "location_name", "sex_id", "sex_name", "age_id", "age_name",
            "cause_id", "cause_name", "metric_id", "metric_name", "year", "val", "upper", "lower"]
    rows = export_rows(**kw) + (extra_rows or [])
    if german:
        rows = [{**r, "sex_name": GERMAN.get(r["sex_name"], r["sex_name"]), "age_name": GERMAN.get(r["age_name"], r["age_name"]),
                 "metric_name": GERMAN.get(r["metric_name"], r["metric_name"]), "measure_name": GERMAN.get(r["measure_name"], r["measure_name"]),
                 "cause_name": GERMAN.get(r["cause_name"], "Ursache " + str(r["cause_id"]))} for r in rows]
    if population_groups:
        cols = ["population_group_id", "population_group_name"] + cols
        rows = [{**r, "population_group_id": g, "population_group_name": ("Gesamtbevölkerung" if german else "All Population") if g == 1 else f"Group {g}",
                 "val": r["val"] if g == 1 or not isinstance(r["val"], (int, float)) else 999_999 - r["val"]}
                for g in population_groups for r in rows]
    if not with_ids:
        cols = [c for c in cols if not c.endswith("_id")]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore", delimiter=delimiter)
        w.writeheader()
        w.writerows(rows)
    return path


def rank_args(export: Path, hierarchy: Path) -> list[str]:
    return ["rank", "--export", str(export), "--hierarchy", str(hierarchy), "--author", "Test Author",
            "--export-date", "2026-10-07", "--permalink", "https://example.org/permalink"]


def rank(gbd, work: Path, **kw) -> int:
    work.mkdir(parents=True, exist_ok=True)
    export = write_export(work / EXPORT_NAME, **kw)
    hierarchy = write_hierarchy(work / HIERARCHY_NAME)
    return gbd.main(rank_args(export, hierarchy))


def fill_sheet(path: Path, decisions: dict[str, tuple[str, str]] | None = None, default: str = "agree", out: Path | None = None,
               comment: str = "") -> Path:
    """Fill in DECISION, ALTERNATIVE_EXPRESSION and COMMENT like a rater would do in Excel.
    Rows without a proposed expression get rule (d) unless a decision is given."""
    decisions = decisions or {}
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    header_row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "rank")
    header = [c.value for c in ws[header_row]]
    col = {name: header.index(name) + 1 for name in ("gbd_cause", "proposed_expression", "DECISION", "ALTERNATIVE_EXPRESSION", "COMMENT")}
    for r in range(header_row + 1, ws.max_row + 1):
        cause = ws.cell(row=r, column=col["gbd_cause"]).value
        if not cause:
            continue
        proposed = ws.cell(row=r, column=col["proposed_expression"]).value or ""
        fallback = default if (proposed and not proposed.startswith("(none")) or default != "agree" else RULE_D
        decision, alternative = decisions.get(cause, (fallback, ""))
        ws.cell(row=r, column=col["DECISION"]).value = decision or None
        ws.cell(row=r, column=col["ALTERNATIVE_EXPRESSION"]).value = alternative or None
        ws.cell(row=r, column=col["COMMENT"]).value = comment or None
    target = out or path
    wb.save(target)
    return target


def edit_cell(path: Path, cause: str, column: str, value) -> None:
    """Change one cell of a sheet outside the helper above (what a careless user might do)."""
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    header_row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "rank")
    header = [c.value for c in ws[header_row]]
    cause_col, target_col = header.index("gbd_cause") + 1, header.index(column) + 1
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.cell(row=r, column=cause_col).value == cause:
            ws.cell(row=r, column=target_col).value = value
            wb.save(path)
            return
    raise AssertionError(f"cause {cause!r} not in {path}")


def sheets_of(gbd) -> tuple[Path, Path]:
    return gbd.RATING_DIR / "gbd_mapping_rater_WP.xlsx", gbd.RATING_DIR / "gbd_mapping_rater_IE.xlsx"


def prepare_sheets(gbd, work: Path) -> tuple[Path, Path]:
    assert rank(gbd, work) == 0
    assert gbd.main(["sheets"]) == 0
    return sheets_of(gbd)


def final_panel(tmp_path: Path) -> dict:
    """Run the four steps on the synthetic export into ``tmp_path`` (laid out like a repository)
    and return the paths: a FINAL, non-fixture panel that agrees with its mapping table, ranking
    and export record, and the amendment text that lists it. Used where the identity check has to
    be exercised end to end."""
    gbd = load_script(tmp_path)
    a, b = prepare_sheets(gbd, tmp_path / "gbd_local")
    fill_sheet(a)
    fill_sheet(b)
    assert gbd.main(["reconcile"]) == 0
    section4 = (FIXTURES / "amendment_v3_0_excerpt.md").read_text(encoding="utf-8").split("## 5. GBD panel")[0]
    amendment = tmp_path / "docs" / "AMENDMENT_v3_0.md"
    amendment.parent.mkdir(parents=True, exist_ok=True)
    amendment.write_text(section4 + "\n## 5. GBD panel\n\n" + gbd.AMENDMENT_TABLE.read_text(encoding="utf-8"), encoding="utf-8")
    return {"gbd": gbd, "root": tmp_path, "panel": gbd.PANEL_YAML, "amendment": amendment, "mapping": gbd.MAPPING,
            "ranking": gbd.RANKING_PUBLIC, "record": gbd.EXPORT_RECORD, "table": gbd.AMENDMENT_TABLE, "sheets": (a, b)}
