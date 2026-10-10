"""scripts/gbd_panel.py: ranking of the export, rater sheets, reconciliation and panel build.

Everything runs offline on synthetic files written into a temporary directory (tests/gbd_helpers.py).
The DALY values are invented for the test and are not GBD estimates."""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

import pytest
import yaml

from airct_benchmark import gbd as G
from airct_benchmark.catalogue import CatalogueError, load_catalogue
from airct_benchmark.queries import build_count_queries
from conftest import FIXTURE_AMENDMENT, FIXTURE_PANEL, FIXTURES
from gbd_helpers import (EXPORT_NAME, HIERARCHY_NAME, LEVEL3, RULE_D, edit_cell, export_rows, fill_sheet, final_panel, hierarchy_rows,
                         level3_total, load_script, prepare_sheets, rank, rank_args, write_export, write_hierarchy)

LRI = "Lower respiratory infections"
NO_PROPOSAL = "Synthetic cause without proposal"


@pytest.fixture
def gbd(tmp_path):
    """The script with every output path redirected into a temporary directory."""
    return load_script(tmp_path)


@pytest.fixture
def work(tmp_path):
    """Folder of the raw files, outside the (temporary) repository tree."""
    return tmp_path / "gbd_local"


# ---- rank ---------------------------------------------------------------------------------------

def test_rank_filters_level_3_and_applies_rules_a_to_c(gbd, work, capsys):
    assert rank(gbd, work) == 0
    out = capsys.readouterr().out
    rows = list(csv.DictReader(open(gbd.RANKING_PUBLIC, encoding="utf-8")))
    assert [r["cause_name"] for r in rows] == [n for n, _ in LEVEL3]          # Level 3 only, in descending order of the value
    assert [int(r["rank"]) for r in rows] == list(range(1, len(LEVEL3) + 1))   # ranks among ALL Level 3 causes, never renumbered
    by_name = {r["cause_name"]: r for r in rows}
    assert by_name["Road injuries"]["exclusion_reason"].startswith("rule (a)")
    assert by_name["Self-harm"]["panel_rule_result"] == "excluded" and by_name["Self-harm"]["level1_group"] == "Injuries"
    assert by_name["Other COVID-19 pandemic-related outcomes"]["exclusion_reason"].startswith("rule (b)")
    assert by_name["Other musculoskeletal disorders"]["exclusion_reason"].startswith("rule (c)")
    assert by_name["Stroke"]["panel_rule_result"] == "kept for mapping" and by_name["Stroke"]["level1_group"] == "NCD"
    assert by_name["Malaria"]["level1_group"] == "CMNN"
    assert sum(r["panel_rule_result"] == "kept for mapping" for r in rows) == len(LEVEL3) - 5

    # the deposited ranking holds no DALY value; the ranking with values is written next to the raw export
    assert "daly_count" not in rows[0] and "daly" not in gbd.RANKING_PUBLIC.read_text(encoding="utf-8").lower()
    full_path = work / gbd.RANKING_FULL_NAME
    full = list(csv.DictReader(open(full_path, encoding="utf-8")))
    assert float(full[0]["daly_count"]) == len(LEVEL3) * 1000
    assert not list(gbd.GBD_DIR.rglob(gbd.RANKING_FULL_NAME))

    record = json.loads(gbd.EXPORT_RECORD.read_text(encoding="utf-8"))
    assert record["settings"] == {"measure": "DALYs (Disability-Adjusted Life Years)", "metric": "Number", "location": "Global",
                                  "age": "All ages", "sex": "Both", "year": 2023, "cause": "all causes or all Level 3 causes"}
    assert record["exported_by"] == "Test Author" and record["export_date"] == "2026-10-07"
    assert record["permalink"] == "https://example.org/permalink"
    assert record["raw_export"]["sha256"] == gbd.sha256_file(work / EXPORT_NAME)
    assert record["cause_hierarchy"]["sha256"] == gbd.sha256_file(work / HIERARCHY_NAME)
    assert record["cause_hierarchy"]["read_from"].endswith("worksheet 'Cause Hierarchy'")
    assert record["ranking_public"]["sha256"] == gbd.sha256_text(gbd.RANKING_PUBLIC)
    assert record["ranking_with_values_kept_by_the_authors"] == {"file_name": gbd.RANKING_FULL_NAME, "sha256": gbd.sha256_file(full_path)}
    assert record["all_causes_check"] == {"done": True, "level3_sum_over_all_causes": 1.0, "tolerance": 0.005}
    assert record["level3_causes_ranked"] == len(LEVEL3) and record["export_aggregates_outside_hierarchy_ignored"] == []
    assert "Used with permission" in record["source_identifier"]
    assert "kept after rules (a) to (c)" in out and "add up to 1.0000 times 'All causes'" in out
    assert "<- excluded, rule (a)" in out


def test_rank_matches_by_name_when_the_export_has_no_ids_and_tolerates_british_spelling(gbd, work):
    spelling = {"Ischemic heart disease": "Ischaemic heart disease", "Diarrheal diseases": "Diarrhoeal diseases"}
    assert rank(gbd, work, with_ids=False, spelling=spelling) == 0
    rows = list(csv.DictReader(open(gbd.RANKING_PUBLIC, encoding="utf-8")))
    # the cause is found by its name; the name written is the one of the hierarchy file
    assert rows[0]["cause_name"] == "Ischemic heart disease" and rows[0]["cause_id"] == "2000"
    assert rows[9]["cause_name"] == "Diarrheal diseases"


def test_rank_reads_an_export_made_with_the_german_interface_by_its_ids(gbd, work, capsys):
    """The Results Tool can be used in several languages; the CSV then names sex, age, metric and every cause in that
    language. The ids are the same in every language, and the GBD 2023 tool adds the population group. The ranking
    must be the one of the English export, with the English cause names of the hierarchy file."""
    work.mkdir()
    citation = work / "citation.txt"
    citation.write_text("GBD Results tool:\nUse the following to cite data included in this download:\nGlobal Burden of Disease "
                        "Collaborative Network.\nGlobal Burden of Disease Study 2023 (GBD 2023) Results.\n", encoding="utf-8")
    english = rank_args(write_export(work / "en.csv"), write_hierarchy(work / HIERARCHY_NAME))
    assert gbd.main(english) == 0
    reference = gbd.RANKING_PUBLIC.read_text(encoding="utf-8")
    capsys.readouterr()

    german = write_export(work / "de.csv", german=True, population_groups=(1, 7))
    assert "Beide" in german.read_text(encoding="utf-8") and "Stroke" not in german.read_text(encoding="utf-8")
    assert gbd.main(rank_args(german, work / HIERARCHY_NAME) + ["--citation-file", str(citation)]) == 0
    out = capsys.readouterr().out
    assert gbd.RANKING_PUBLIC.read_text(encoding="utf-8") == reference          # same ranks, same English names, same rule results
    assert "the export is not in English (sex 'Beide', age 'Alle Altersgruppen', metric 'Anzahl')" in out
    assert "add up to 1.0000 times 'All causes'" in out and "Ischemic heart disease" in out
    record = json.loads(gbd.EXPORT_RECORD.read_text(encoding="utf-8"))
    assert record["stratum_as_named_in_the_export"] == {"measure": "DALYs (Disability-Adjusted Life Years)", "location": "Global",
                                                        "sex": "Beide", "age": "Alle Altersgruppen", "metric": "Anzahl"}
    assert record["data_citation_delivered_with_the_download"].startswith("GBD Results tool: Use the following to cite")
    assert record["names_spelled_differently_in_export_and_hierarchy"] == [] and record["all_causes_check"]["done"] is True
    assert any("cause name(s) of the export differ from the names of the hierarchy file" in n for n in record["notes"])

    # without ids a German export cannot be read, and the message says what to do
    capsys.readouterr()
    assert gbd.main(rank_args(write_export(work / "de_names.csv", german=True, with_ids=False), work / HIERARCHY_NAME)) == 2
    err = capsys.readouterr().err
    assert "no row matches Measure DALYs" in err and "download it with names and ids" in err

    # the measure Deaths instead of DALYs, as the German interface writes it: the message names the wrong setting
    deaths = work / "deaths.csv"
    deaths.write_text(german.read_text(encoding="utf-8").replace(",2,DALYs (Disability-Adjusted Life Years),", ",1,Todesfälle,"), encoding="utf-8")
    capsys.readouterr()
    assert gbd.main(rank_args(deaths, work / HIERARCHY_NAME)) == 2
    err = capsys.readouterr().err
    assert "Wrong setting: measure must be DALYs (id 2), the file holds ['1 / Todesfälle']" in err and "sex must be" not in err


def test_rank_accepts_the_level_3_causes_alone_with_the_aggregates_of_the_real_hierarchy(gbd, work, capsys):
    """The GBD 2023 hierarchy file lists aggregates such as 'Total cancers' at Level 1. They are never ranked and
    must not stop an export that holds the Level 3 causes only; the missing comparison with the total is noted."""
    work.mkdir()
    hierarchy = write_hierarchy(work / "h.csv", extra=[
        {"Cause ID": "1029", "Cause Name": "Total cancers", "Parent ID": "294", "Level": "1"},
        {"Cause ID": "1026", "Cause Name": "Total burden related to hepatitis B", "Parent ID": "294", "Level": "1"}])
    others = tuple(r["Cause Name"] for r in hierarchy_rows() if r["Level"] != "3")
    export = write_export(work / "e.csv", drop=others, german=True, population_groups=(1,))
    assert gbd.main(rank_args(export, hierarchy)) == 0
    out = capsys.readouterr().out
    assert "could not be compared with 'All causes'" in out and "were not ranked: ['Total burden related to hepatitis B', 'Total cancers']" in out
    rows = list(csv.DictReader(open(gbd.RANKING_PUBLIC, encoding="utf-8")))
    assert [r["cause_name"] for r in rows] == [n for n, _ in LEVEL3]
    record = json.loads(gbd.EXPORT_RECORD.read_text(encoding="utf-8"))
    assert record["all_causes_check"]["done"] is False and len(record["hierarchy_aggregates_not_ranked"]) == 2


def test_rank_names_the_export_that_holds_only_all_causes(gbd, work, capsys):
    """The default of the Results Tool is the single cause 'All causes'. Such an export has one row."""
    work.mkdir()
    export = work / "only_all_causes.csv"
    export.write_text("population_group_id,population_group_name,measure_id,measure_name,location_id,location_name,sex_id,sex_name,"
                      "age_id,age_name,cause_id,cause_name,metric_id,metric_name,year,val,upper,lower\n"
                      "1,Gesamtbevölkerung,2,DALYs (Disability-Adjusted Life Years),1,Global,3,Beide,22,Alle Altersgruppen,294,"
                      "Alle Ursachen,1,Anzahl,2023,5000.5,6000.5,4000.5\n", encoding="utf-8")
    assert gbd.main(rank_args(export, write_hierarchy(work / "h.csv"))) == 2
    err = capsys.readouterr().err
    assert "holds no Level 3 cause at all (1 row(s) of the stratum: ['All causes'])" in err
    assert "the field 'Cause' was left at its default" in err and not gbd.RANKING_PUBLIC.exists()


def test_rank_requires_who_when_and_where(gbd, work, capsys):
    export, hierarchy = write_export(work.parent / "e.csv"), write_hierarchy(work.parent / "h.csv")
    with pytest.raises(SystemExit):                                         # argparse: the three options are required
        gbd.main(["rank", "--export", str(export), "--hierarchy", str(hierarchy)])
    capsys.readouterr()
    args = rank_args(export, hierarchy)
    assert gbd.main([a if a != "2026-10-07" else "07.10.2026" for a in args]) == 2
    assert "YYYY-MM-DD" in capsys.readouterr().err
    assert gbd.main([a if not a.startswith("https://") else "see mail" for a in args]) == 2
    assert not gbd.RANKING_PUBLIC.exists()


@pytest.mark.parametrize("case, message", [
    ("missing", "2 Level 3 cause(s) of the hierarchy are missing from the export"),
    ("foreign_unchecked", "1 cause(s) of the export are not in the hierarchy file"),
    ("other_round", "is 'Cerebrovascular disease' in the export but 'Stroke' in the hierarchy"),
    ("incomplete_sum", "add up to"),
    ("tie", "exactly the same value"),
    ("negative", "not a finite, non-negative number"),
    ("decimal_comma", "decimal comma"),
    ("semicolons", "separated by semicolons"),
    ("zip", "is a ZIP archive"),
    ("risk_export", "risk-factor columns"),
    ("wrong_stratum", "no row matches Measure DALYs"),
    ("no_parent", "'Parent ID' and 'Level'"),
    ("level1_renamed", "Level 1 group"),
    ("total_unchecked", "lists aggregates across causes among the causes of Level 3"),
    ("leaf_kept", "is a cause of Level 2 without Level 3 causes below it and is not excluded"),
    ("leaf_without_value_unchecked", "have no row in the export"),
    ("rules_differ_by_name", "give different results for the two names"),
    ("duplicate", "occurs twice after filtering"),
])
def test_rank_stops_with_a_clear_message_instead_of_a_wrong_ranking(gbd, work, capsys, case, message):
    work.mkdir()
    export, hierarchy = work / "export.csv", work / "hierarchy.csv"
    write_hierarchy(hierarchy)
    extra = {"measure_id": 2, "measure_name": "DALYs (Disability-Adjusted Life Years)", "location_id": 1, "location_name": "Global",
             "sex_id": 3, "sex_name": "Both", "age_id": 22, "age_name": "All ages", "metric_id": 1, "metric_name": "Number", "year": 2023,
             "upper": 1, "lower": 1}
    if case == "missing":
        write_export(export, drop=("Malaria", "Tuberculosis"))
    elif case == "foreign_unchecked":          # no row 'All causes': the total cannot vouch for the ranking
        write_export(export, drop=("All causes",), extra_rows=[{**extra, "cause_id": 9998, "cause_name": "A cause of another GBD round", "val": 5}])
    elif case == "other_round":
        write_export(export, spelling={"Stroke": "Cerebrovascular disease"})
    elif case == "incomplete_sum":
        write_export(export, values={"All causes": 5_000_000})
    elif case == "tie":
        values = {"Malaria": 21000}                                           # the value of the next cause in the list
        write_export(export, values={**values, "All causes": level3_total(values)})
    elif case == "negative":
        write_export(export, values={"Malaria": -5, "All causes": level3_total({"Malaria": -5})})
    elif case == "decimal_comma":
        write_export(export, values={"Malaria": "22000,5"})
    elif case == "semicolons":
        write_export(export, delimiter=";")
    elif case == "zip":
        shutil.make_archive(str(work / "download"), "zip", root_dir=work, base_dir="hierarchy.csv")
        export = work / "download.zip"
    elif case == "risk_export":
        write_export(export)
        text = export.read_text(encoding="utf-8").replace("cause_id,cause_name", "cause_id,cause_name,rei_id,rei_name", 1)
        export.write_text(text, encoding="utf-8")
    elif case == "wrong_stratum":
        export.write_text("measure_name,location_name,sex_name,age_name,cause_name,metric_name,year,val\n"
                          "Deaths,Global,Both,All ages,Stroke,Number,2023,5\n", encoding="utf-8")
    elif case == "no_parent":
        write_export(export)
        write_hierarchy(hierarchy, columns=("Cause ID", "Cause Name", "Level"))
    elif case == "level1_renamed":
        write_hierarchy(hierarchy, rename={"Injuries": "C. Injuries"})
        write_export(export, spelling={"Injuries": "C. Injuries"})
    elif case == "total_unchecked":
        write_export(export, drop=("All causes",), spelling={"Asthma": "Total burden of something"})
        write_hierarchy(hierarchy, rename={"Asthma": "Total burden of something"})
    elif case == "leaf_kept":                  # a Level 2 cause without Level 3 causes that no rule excludes: a protocol decision
        write_hierarchy(hierarchy, extra=[{"Cause ID": "960", "Cause Name": "A new cause group", "Parent ID": "295", "Level": "2"}])
        write_export(export, values={"All causes": level3_total() + 700}, extra_rows=[{**extra, "cause_id": 960, "cause_name": "A new cause group", "val": 700}])
    elif case == "leaf_without_value_unchecked":
        write_hierarchy(hierarchy, extra=[{"Cause ID": "960", "Cause Name": "Other pandemic-related outcomes", "Parent ID": "295", "Level": "2"}])
        write_export(export, drop=("All causes",))
    elif case == "rules_differ_by_name":       # same id, the same name in another spelling, but rule (c) reads the two differently
        write_export(export, spelling={"Other musculoskeletal disorders": "Other-musculoskeletal disorders"})
    elif case == "duplicate":
        write_export(export, extra_rows=[r for r in export_rows() if r["cause_name"] == "Malaria"][:1])
    assert gbd.main(rank_args(export, hierarchy)) == 2
    err = capsys.readouterr().err
    assert err.startswith("\nSTOPPED:") and message in err and "Traceback" not in err
    assert not gbd.RANKING_PUBLIC.exists() and not gbd.EXPORT_RECORD.exists()


def test_rank_tolerates_what_the_total_vouches_for_and_records_it(gbd, work, capsys):
    """With a row 'All causes' that the ranked causes add up to, the order of the leading causes cannot be wrong.
    Then a cause of the export that the hierarchy does not know, an aggregate listed in the hierarchy, a name
    spelled differently and a cause of Level 2 without Level 3 causes (excluded by a rule) are reported and
    recorded, and the ranking is made."""
    work.mkdir()
    extra = {"measure_id": 2, "measure_name": "DALYs (Disability-Adjusted Life Years)", "location_id": 1, "location_name": "Global",
             "sex_id": 3, "sex_name": "Both", "age_id": 22, "age_name": "All ages", "metric_id": 1, "metric_name": "Number", "year": 2023,
             "upper": 1, "lower": 1}
    hierarchy = write_hierarchy(work / "h.csv", extra=[
        {"Cause ID": "960", "Cause Name": "Other pandemic-related outcomes", "Parent ID": "295", "Level": "2"},      # no Level 3 cause below it
        {"Cause ID": "970", "Cause Name": "Total burden related to something", "Parent ID": "409", "Level": "3"},     # an aggregate
    ])
    export = write_export(work / "e.csv", values={"All causes": level3_total() + 1500},
                          spelling={"Alzheimer's disease and other dementias": "Alzheimer disease and other dementias"},
                          extra_rows=[{**extra, "cause_id": 960, "cause_name": "Other pandemic-related outcomes", "val": 1500},
                                      {**extra, "cause_id": 970, "cause_name": "Total burden related to something", "val": 77777},
                                      {**extra, "cause_id": 9998, "cause_name": "A sub-cause the hierarchy file does not list", "val": 5}])
    assert gbd.main(rank_args(export, hierarchy)) == 0
    out = capsys.readouterr().out
    assert f"{len(LEVEL3) + 1} causes ranked by global DALYs 2023 ({len(LEVEL3)} of Level 3, 1 of Level 1 or 2 without causes below them)" in out
    assert "add up to 1.0000 times 'All causes'" in out
    assert "1 cause(s) of the export are not in the hierarchy file and were not ranked" in out
    assert "aggregates across causes listed in the hierarchy file at Level 1 to 3 were not ranked" in out
    assert "spelling differs, treated as the same cause: cause id" in out and "ranked with the Level 3 causes, as GBD reports them" in out
    rows = {r["cause_name"]: r for r in csv.DictReader(open(gbd.RANKING_PUBLIC, encoding="utf-8"))}
    assert len(rows) == len(LEVEL3) + 1 and "Total burden related to something" not in rows
    leaf = rows["Other pandemic-related outcomes"]
    assert leaf["panel_rule_result"] == "excluded" and leaf["exclusion_reason"].startswith("rule (b): pandemic-related outcome category; cause of Level 2")
    assert rows["Alzheimer's disease and other dementias"]["panel_rule_result"] == "kept for mapping"    # the name of the hierarchy file
    assert [r for r in rows.values() if r["panel_rule_result"] == "kept for mapping"][0]["cause_name"] == "Ischemic heart disease"
    record = json.loads(gbd.EXPORT_RECORD.read_text(encoding="utf-8"))
    assert record["all_causes_check"]["done"] and record["export_causes_outside_hierarchy_not_ranked"] == ["A sub-cause the hierarchy file does not list"]
    assert record["hierarchy_aggregates_not_ranked"] == ["Total burden related to something"]
    assert record["causes_above_level_3_ranked_as_themselves"] == ["Other pandemic-related outcomes"]
    assert len(record["names_spelled_differently_in_export_and_hierarchy"]) == 1


def test_rank_reads_all_causes_from_the_export_when_the_hierarchy_has_no_level_0(gbd, work, capsys):
    work.mkdir()
    hierarchy = write_hierarchy(work / "h.csv", drop=("All causes",))
    assert gbd.main(rank_args(write_export(work / "e.csv"), hierarchy)) == 0
    assert "add up to 1.0000 times 'All causes'" in capsys.readouterr().out
    assert json.loads(gbd.EXPORT_RECORD.read_text(encoding="utf-8"))["all_causes_check"]["done"] is True
    capsys.readouterr()
    assert gbd.main(rank_args(write_export(work / "e2.csv", values={"All causes": 5_000_000}), hierarchy)) == 2
    assert "add up to" in capsys.readouterr().err


def test_rank_ignores_and_records_aggregates_across_causes(gbd, work, capsys):
    work.mkdir()
    total = {"measure_id": 2, "measure_name": "DALYs (Disability-Adjusted Life Years)", "location_id": 1, "location_name": "Global",
             "sex_id": 3, "sex_name": "Both", "age_id": 22, "age_name": "All ages", "cause_id": 9999, "cause_name": "Total cancers",
             "metric_id": 1, "metric_name": "Number", "year": 2023, "val": 99000000, "upper": 1, "lower": 1}
    export, hierarchy = write_export(work / "e.csv", extra_rows=[total]), write_hierarchy(work / "h.csv")
    assert gbd.main(rank_args(export, hierarchy)) == 0
    assert "1 aggregate(s) across causes" in capsys.readouterr().out
    rows = list(csv.DictReader(open(gbd.RANKING_PUBLIC, encoding="utf-8")))
    assert rows[0]["cause_name"] == "Ischemic heart disease" and len(rows) == len(LEVEL3)
    assert json.loads(gbd.EXPORT_RECORD.read_text(encoding="utf-8"))["export_aggregates_outside_hierarchy_ignored"] == ["Total cancers"]


def test_rank_warns_when_the_raw_files_lie_inside_the_repository(gbd, tmp_path, capsys):
    inside = gbd.GBD_DIR / "local"                                          # GBD_DIR is redirected below tmp_path; REPO is not
    gbd.REPO = tmp_path
    assert rank(gbd, inside) == 0
    out = capsys.readouterr().out
    assert "these files lie inside the repository folder" in out and "must not be published" in out


# ---- sheets ---------------------------------------------------------------------------------------

def test_sheets_are_prefilled_and_a_filled_sheet_is_never_overwritten(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    out = capsys.readouterr().out
    assert f"no proposal in the candidates table for: ['{NO_PROPOSAL}']" in out
    rater, rows = gbd.read_sheet(a)
    assert rater == "WP" and len(rows) == 30
    causes = [r["gbd_cause"] for r in sorted(rows.values(), key=lambda r: int(r["rank"]))]
    assert causes[:6] == ["Ischemic heart disease", "Neonatal disorders", "Stroke", LRI, "Diabetes mellitus",
                          "Chronic obstructive pulmonary disease"]                    # Road injuries is not rated
    assert "Road injuries" not in causes and "Other musculoskeletal disorders" not in causes
    stroke = rows[gbd.norm_name("Stroke")]
    assert stroke["proposed_expression"] == '"Stroke"[Mesh]' and stroke["mesh_ui"] == "D020521"
    assert stroke["mesh_browser_link"] == "https://meshb.nlm.nih.gov/record/ui?ui=D020521"
    assert stroke["DECISION"] == "" and stroke["ALTERNATIVE_EXPRESSION"] == ""
    assert rows[gbd.norm_name("Hypertensive heart disease")]["proposed_expression"] == "(none proposed: rule d)"
    assert rows[gbd.norm_name(NO_PROPOSAL)]["inclusion_note"].startswith("NO PROPOSAL")
    lri = rows[gbd.norm_name(LRI)]
    assert lri["proposed_expression"] == '("Pneumonia"[Mesh] NOT "COVID-19"[Mesh])' and '"Pneumonia"[Mesh] (plain' in lri["alternatives"]

    # a second call must not destroy a sheet that is already filled in
    fill_sheet(a, comment="my reason")
    assert gbd.main(["sheets"]) == 2
    assert "exist already" in capsys.readouterr().err
    assert gbd.read_sheet(a)[1][gbd.norm_name("Malaria")] ["COMMENT"] == "my reason"
    # more causes without losing decisions: new folder, decisions carried over
    more = gbd.GBD_DIR / "rating_2"
    assert gbd.main(["sheets", "--n", "31", "--out-dir", str(more), "--carry-from", str(gbd.RATING_DIR)]) == 0
    _, carried = gbd.read_sheet(more / "gbd_mapping_rater_WP.xlsx")
    assert len(carried) == 31 and carried[gbd.norm_name("Malaria")]["DECISION"] == "agree" and carried[gbd.norm_name("Malaria")]["COMMENT"] == "my reason"
    assert gbd.read_sheet(more / "gbd_mapping_rater_IE.xlsx")[1][gbd.norm_name("Malaria")]["DECISION"] == ""
    assert gbd.main(["sheets", "--raters", "WP", "WP"]) == 2 and gbd.main(["sheets", "--raters", "W/P", "IE"]) == 2
    assert gbd.main(["sheets", "--n", "10"]) == 2


def test_sheet_dropdown_rejects_free_text_and_links_are_clickable(gbd, work):
    import openpyxl

    a, _ = prepare_sheets(gbd, work)
    ws = openpyxl.load_workbook(a)["Mapping"]
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1 and validations[0].showErrorMessage and "agree" in validations[0].formula1
    header_row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "rank")
    link = ws.cell(row=header_row + 1, column=5)
    assert link.hyperlink is not None and link.hyperlink.target == link.value == "https://meshb.nlm.nih.gov/record/ui?ui=D017202"


# ---- reconcile: decisions -----------------------------------------------------------------------------

def test_reconcile_writes_the_final_panel_when_the_raters_agree(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a, comment="checked in the MeSH Browser")
    fill_sheet(b)
    capsys.readouterr()
    assert gbd.main(["reconcile"]) == 0                                     # the two sheets are found in data/gbd/rating
    out = capsys.readouterr().out
    assert "30 causes reconciled: 30 agreed independently, 0 after discussion." in out
    assert "status final, 20 members, ranks 1 to 23" in out

    mapping = list(csv.DictReader(open(gbd.MAPPING, encoding="utf-8")))
    by_cause = {r["gbd_cause"]: r for r in mapping}
    assert len(mapping) == 30 and by_cause["Stroke"]["cause_id"] == "2002" and by_cause["Stroke"]["level1_group"] == "NCD"
    assert by_cause["Hypertensive heart disease"]["decision"] == "excluded rule (d)"
    assert by_cause[NO_PROPOSAL]["decision"] == "excluded rule (d)" and by_cause[NO_PROPOSAL]["choice_WP"] == "rule (d)"
    assert by_cause["Malaria"]["comment_WP"] == "checked in the MeSH Browser" and by_cause["Malaria"]["reconciliation"] == "agreed independently"

    panel = yaml.safe_load(gbd.PANEL_YAML.read_text(encoding="utf-8"))
    assert panel["status"] == "final" and "test_fixture" not in panel and panel["source"]["raters"] == "IE and WP"
    assert panel["source"]["mapping_table"] == "data/gbd/gbd_panel_mesh_mapping.csv"
    assert panel["source"]["mapping_table_sha256"] == gbd.sha256_text(gbd.MAPPING)
    assert panel["source"]["ranking_sha256"] == gbd.sha256_text(gbd.RANKING_PUBLIC)
    assert panel["source"]["raw_export_sha256"] == gbd.sha256_file(work / EXPORT_NAME)
    assert panel["source"]["candidates_sha256"] == gbd.sha256_text(gbd.CANDIDATES)
    assert [m["rank"] for m in panel["members"]] == [1, 2, 3, 4, 5, 7, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 20, 21, 22, 23]
    assert [m["shared_with"] for m in panel["members"] if m["shared_with"]] == ["STROKE", "DM"]

    # the pipeline accepts the panel; the hash stated in the amendment table is the one of the frozen files
    cat = load_catalogue(gbd_panel_path=gbd.PANEL_YAML)
    assert cat.gbd_is_final and not cat.gbd_is_test_fixture and len(build_count_queries(cat)) == 4326
    table = gbd.AMENDMENT_TABLE.read_text(encoding="utf-8")
    assert f"of config/gbd_panel.yaml {cat.gbd_sha256()}" in table and gbd.sha256_text(gbd.PANEL_YAML) == cat.gbd_sha256()
    assert "Road injuries (rank 6, rule (a): Level 1 group Injuries)" in table
    assert "Other musculoskeletal disorders (rank 19, rule (c): residual cause (name begins with Other))" in table
    assert "by Test Author, permalink https://example.org/permalink" in table
    assert "Used with permission" in table and "4,326 per run" in table and "mapped by the rule of 5.3 by IE and WP" in table
    assert G.table_rows_in_text(table) == [G.table_row(m) for m in panel["members"]]
    assert G.panel_file_problems(panel, gbd.PANEL_YAML.parent.parent, {'"Stroke"[Mesh]': "STROKE", '"Diabetes Mellitus"[Mesh]': "DM"}) == []


def test_disagreements_go_through_a_consensus_sheet(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a, {LRI: ("use alternative", '"Pneumonia"[Mesh]')}, comment="plain heading is simpler")
    fill_sheet(b)
    capsys.readouterr()
    assert gbd.main(["reconcile"]) == 1
    out = capsys.readouterr().out
    assert "1 disagreement(s) between IE and WP. No panel was written." in out
    assert 'WP: "Pneumonia"[Mesh]   (plain heading is simpler)' in out and 'IE: ("Pneumonia"[Mesh] NOT "COVID-19"[Mesh])' in out
    assert not gbd.MAPPING.exists() and not gbd.PANEL_YAML.exists()
    consensus = gbd.RATING_DIR / "gbd_mapping_consensus.xlsx"
    rater, rows = gbd.read_sheet(consensus)
    assert rater == "CONSENSUS" and list(rows) == [gbd.norm_name(LRI)]
    assert rows[gbd.norm_name(LRI)]["choice_WP"] == '"Pneumonia"[Mesh]' and rows[gbd.norm_name(LRI)]["DECISION"] == ""

    # the consensus sheet is still empty: the same answer, and it is not overwritten
    assert gbd.main(["reconcile"]) == 1
    # a rater sheet cannot serve as the consensus
    assert gbd.main(["reconcile", "--consensus", str(a)]) == 2
    assert "is not a consensus sheet" in capsys.readouterr().err
    # the agreed decision is entered, the same command is run again
    fill_sheet(consensus, {LRI: ("agree", "")}, comment="COVID-19 is counted separately by GBD")
    assert gbd.main(["reconcile"]) == 0
    out = capsys.readouterr().out
    assert "resolved by the consensus sheet" in out and "29 agreed independently, 1 after discussion" in out
    row = {r["gbd_cause"]: r for r in csv.DictReader(open(gbd.MAPPING, encoding="utf-8"))}[LRI]
    assert row["reconciliation"] == "agreed after discussion" and row["choice_WP"] == '"Pneumonia"[Mesh]'
    assert row["expression"] == '("Pneumonia"[Mesh] NOT "COVID-19"[Mesh])' and row["comment_consensus"] == "COVID-19 is counted separately by GBD"

    # a consensus sheet that no longer fits the rater sheets is refused
    fill_sheet(a, {LRI: ("use alternative", '"Respiratory Tract Infections"[Mesh]')})
    assert gbd.main(["reconcile"]) == 2
    assert "delete the consensus sheet" in capsys.readouterr().err


def test_both_raters_can_agree_on_a_listed_alternative_or_on_rule_d(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    decisions = {LRI: ("use alternative", '"Pneumonia"[Mesh]'), "Low back pain": (RULE_D, "")}
    fill_sheet(a, decisions)
    fill_sheet(b, decisions)
    assert gbd.main(["reconcile"]) == 0
    panel = yaml.safe_load(gbd.PANEL_YAML.read_text(encoding="utf-8"))
    members = {m["cause"]: m for m in panel["members"]}
    assert members[LRI]["expression"] == '"Pneumonia"[Mesh]' and members[LRI]["mesh_ui"] == "D011014"
    assert "COVID-19 literature from 2020 on" in members[LRI]["inclusion_note"]      # the facts of the alternative, not of the proposal
    causes = [m["cause"] for m in panel["members"]]
    assert "Low back pain" not in causes and causes[-1] == "Dietary iron deficiency" and len(causes) == 20
    assert "Low back pain (rank 9, rule (d): no single MeSH descriptor)" in gbd.AMENDMENT_TABLE.read_text(encoding="utf-8")


@pytest.mark.parametrize("decisions, message", [
    ({"Malaria": ("", "")}, "DECISION is empty"),
    ({"Malaria": ("agree", '"Malaria, Falciparum"[Mesh]')}, "DECISION is 'agree' but ALTERNATIVE_EXPRESSION is filled in"),
    ({"Malaria": (RULE_D, '"Malaria, Falciparum"[Mesh]')}, "DECISION is 'no single descriptor' but ALTERNATIVE_EXPRESSION is filled in"),
    ({"Malaria": ("use alternative", "Malaria[Mesh]")}, "is not of the form"),
    ({"Malaria": ("use alternative", '"Malaria"[Mesh]')}, "the alternative equals the proposal"),
    ({NO_PROPOSAL: ("agree", "")}, "there is no proposed expression to agree with"),
    ({"Hypertensive heart disease": ("agree", "")}, "there is no proposed expression to agree with"),
    ({"Stroke": ("use alternative", '"Cerebrovascular Disorders"[Mesh]')}, "named field of the study"),
    ({"Diabetes mellitus": (RULE_D, "")}, "named field of the study"),
])
def test_reconcile_guesses_nothing(gbd, work, capsys, decisions, message):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a, decisions)
    fill_sheet(b, decisions)
    capsys.readouterr()
    assert gbd.main(["reconcile"]) == 2
    err = capsys.readouterr().err
    assert message in err and "Traceback" not in err
    assert not gbd.MAPPING.exists() and not gbd.PANEL_YAML.exists()


def test_reconcile_reports_every_unreadable_row_at_once(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a, {"Malaria": ("", ""), "Tuberculosis": ("", "")})
    fill_sheet(b, {"Asthma": ("perhaps", "")})
    capsys.readouterr()
    assert gbd.main(["reconcile"]) == 2
    err = capsys.readouterr().err
    assert "3 row(s) cannot be read as a decision" in err
    assert "sheet of WP, Malaria" in err and "sheet of WP, Tuberculosis" in err and "sheet of IE, Asthma" in err


def test_an_agreed_alternative_needs_a_fact_row_for_that_cause(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    decisions = {"Malaria": ("use alternative", '"Malaria, Falciparum"[Mesh]')}       # not in the candidates table
    fill_sheet(a, decisions)
    fill_sheet(b, decisions)
    assert gbd.main(["reconcile"]) == 2
    assert "has no row for this cause" in capsys.readouterr().err
    # a descriptor that is listed, but for another cause, does not lend its facts
    decisions = {"Malaria": ("use alternative", '"Tuberculosis"[Mesh]')}
    fill_sheet(a, decisions)
    fill_sheet(b, decisions)
    assert gbd.main(["reconcile"]) == 2
    assert "has no row for this cause" in capsys.readouterr().err
    assert not gbd.MAPPING.exists() and not gbd.PANEL_YAML.exists()


# ---- reconcile: the sheets against ranking and candidates ----------------------------------------------

def test_reconcile_refuses_sheets_whose_ranks_were_changed(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a)
    fill_sheet(b)
    edit_cell(a, "Breast cancer", "rank", 20)
    edit_cell(a, "HIV/AIDS", "rank", 31)
    assert gbd.main(["reconcile"]) == 2
    err = capsys.readouterr().err
    assert "do not match the ranking" in err and "'Breast cancer' has rank 20, the ranking says 31" in err
    assert not gbd.PANEL_YAML.exists()


def test_reconcile_refuses_sheets_that_belong_to_another_ranking(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a)
    fill_sheet(b)
    # the export is corrected and ranked again: Asthma is now the second cause
    values = {"Asthma": 35500}
    export = write_export(work / "corrected.csv", values={**values, "All causes": level3_total(values)})
    assert gbd.main(rank_args(export, work / HIERARCHY_NAME)) == 0
    assert "Rater sheets exist already" in capsys.readouterr().out
    assert gbd.main(["reconcile"]) == 2
    err = capsys.readouterr().err
    assert "do not match the ranking" in err and "new sheets are needed" in err
    assert not gbd.PANEL_YAML.exists()


def test_reconcile_refuses_added_or_deleted_rows(gbd, work, capsys):
    import openpyxl

    a, b = prepare_sheets(gbd, work)
    fill_sheet(a)
    fill_sheet(b)
    for path in (a, b):                                                     # both raters add the same excluded cause
        wb = openpyxl.load_workbook(path)
        ws = wb["Mapping"]
        ws.append([6, "Road injuries", '"Accidents, Traffic"[Mesh]', "D000063", "", "1966-01-01", "exact", "", "", "", "", "agree", None, None])
        wb.save(path)
    assert gbd.main(["reconcile"]) == 2
    assert "'Road injuries' is excluded by the ranking" in capsys.readouterr().err
    a2, b2 = prepare_sheets_again(gbd)
    fill_sheet(a2)
    fill_sheet(b2)
    wb = openpyxl.load_workbook(a2)
    ws = wb["Mapping"]
    header_row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "rank")
    assert ws.cell(row=header_row + 3, column=2).value == "Stroke"
    for cell in ws[header_row + 3]:                                         # one rater empties the Stroke row
        cell.value = None
    wb.save(a2)
    assert gbd.main(["reconcile", "--sheets", str(a2), str(b2)]) == 2
    err = capsys.readouterr().err                # the emptied row keeps its link cell, or is gone: refused either way
    assert ("has text but no cause" in err or "do not list the same causes" in err) and not gbd.PANEL_YAML.exists()


def prepare_sheets_again(gbd) -> tuple[Path, Path]:
    assert gbd.main(["sheets", "--force"]) == 0
    return gbd.RATING_DIR / "gbd_mapping_rater_WP.xlsx", gbd.RATING_DIR / "gbd_mapping_rater_IE.xlsx"


def test_reconcile_refuses_when_sheet_and_candidates_table_differ(gbd, work, tmp_path, capsys):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a)
    fill_sheet(b)
    # a rater changes a pre-filled cell
    edit_cell(a, "Malaria", "proposed_expression", '"Malaria, Falciparum"[Mesh]')
    assert gbd.main(["reconcile"]) == 2
    assert "sheet of WP, Malaria: column proposed_expression differs from the candidates table" in capsys.readouterr().err
    edit_cell(a, "Malaria", "proposed_expression", '"Malaria"[Mesh]')
    # the candidates table is changed after the sheets were written
    changed = tmp_path / "candidates_changed.csv"
    changed.write_text(gbd.CANDIDATES.read_text(encoding="utf-8").replace('"""Malaria""[Mesh]",D008288,1960-01-01', '"""Malaria""[Mesh]",D008288,2016-01-01'), encoding="utf-8")
    assert "2016-01-01" in changed.read_text(encoding="utf-8")
    assert gbd.main(["reconcile", "--candidates", str(changed)]) == 2
    err = capsys.readouterr().err
    assert "Malaria: column date_introduced differs from the candidates table" in err and not gbd.PANEL_YAML.exists()


def test_reconcile_and_sheet_reading_tolerate_what_office_programs_do(gbd, work, capsys):
    import openpyxl

    a, b = prepare_sheets(gbd, work)
    fill_sheet(a, {"Malaria": ("  AGREE ", "")})                              # capitals and blanks
    fill_sheet(b, {LRI: ("use alternative", "„Pneumonia“[Mesh]")})   # typographic quotation marks
    fill_sheet(a, {LRI: ("use alternative", '"Pneumonia"[Mesh]'), "Malaria": ("  AGREE ", "")})
    wb = openpyxl.load_workbook(b)                                          # worksheet renamed, rows sorted in reverse
    ws = wb["Mapping"]
    ws.title = "Tabelle1"
    header_row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "rank")
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row)]
    for i, values in enumerate(reversed(data), header_row + 1):
        for j, v in enumerate(values, 1):
            ws.cell(row=i, column=j).value = v
    renamed = b.with_name("Bogen Eyüpoglu final.xlsx")
    wb.save(renamed)
    assert gbd.main(["reconcile", "--sheets", str(a), str(renamed)]) == 0
    assert "30 agreed independently" in capsys.readouterr().out


def test_file_problems_end_with_stopped_not_with_a_traceback(gbd, work, tmp_path, capsys, monkeypatch):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a)
    fill_sheet(b)
    (tmp_path / "not_excel.xlsx").write_text("this is not a workbook", encoding="utf-8")
    assert gbd.main(["reconcile", "--sheets", str(a), str(tmp_path / "not_excel.xlsx")]) == 2
    assert "cannot be read as an Excel file" in capsys.readouterr().err
    assert gbd.main(["reconcile", "--sheets", str(a), str(tmp_path / "missing.xlsx")]) == 2
    assert "sheet not found" in capsys.readouterr().err
    assert gbd.main(["reconcile", "--ranking", str(tmp_path / "no_ranking.csv")]) == 2
    assert gbd.main(["reconcile", "--candidates", str(gbd.RANKING_PUBLIC)]) == 2                 # a wrong file for the candidates
    assert "lacks the column(s)" in capsys.readouterr().err
    assert gbd.main(["build", "--mapping", str(gbd.RANKING_PUBLIC)]) == 2
    assert "lacks the column(s)" in capsys.readouterr().err

    def locked(*args, **kwargs):
        raise PermissionError(13, "Permission denied", str(gbd.MAPPING))

    monkeypatch.setattr(gbd, "write_text_lf", locked)                         # an output file is open in Excel
    assert gbd.main(["reconcile"]) == 2
    err = capsys.readouterr().err
    assert "probably open in Excel" in err and "Traceback" not in err


# ---- build ----------------------------------------------------------------------------------------------

def test_a_rejected_panel_leaves_no_trace(gbd, work, tmp_path, capsys):
    a, b = prepare_sheets(gbd, work)
    # both raters choose, for Malaria, the expression that Tuberculosis already has: two members would share one expression
    candidates = tmp_path / "candidates_plus.csv"
    extra = 'Malaria,alternative,"""Tuberculosis""[Mesh]",D014376,1960-01-01,,exact,,,synthetic row of the test,,,test,\n'
    candidates.write_text(gbd.CANDIDATES.read_text(encoding="utf-8") + extra, encoding="utf-8")
    decisions = {"Malaria": ("use alternative", '"Tuberculosis"[Mesh]')}
    fill_sheet(a, decisions)
    fill_sheet(b, decisions)
    assert gbd.main(["reconcile", "--candidates", str(candidates)]) == 2
    err = capsys.readouterr().err
    assert "the pipeline does not accept this panel, nothing was written" in err and "share one expression" in err
    assert not gbd.PANEL_YAML.exists() and not gbd.MAPPING.exists() and not list(gbd.PANEL_YAML.parent.glob("*.tmp"))


def test_build_is_provisional_by_default_and_final_only_from_a_reconciled_mapping(gbd, work, capsys):
    # a mapping table that 'reconcile' did not write (here: the fixture table of the test suite) gives a provisional panel ...
    placeholder = FIXTURES / "gbd_mapping_fixture.csv"
    assert gbd.main(["build", "--mapping", str(placeholder)]) == 0
    assert yaml.safe_load(gbd.PANEL_YAML.read_text(encoding="utf-8"))["status"] == "provisional"
    assert "PROVISIONAL, not for the OSF update" in gbd.AMENDMENT_TABLE.read_text(encoding="utf-8")
    # ... and never a final one; the provisional file stays as it was
    capsys.readouterr()
    assert gbd.main(["build", "--mapping", str(placeholder), "--status", "final"]) == 2
    assert "not written by 'reconcile'" in capsys.readouterr().err
    assert yaml.safe_load(gbd.PANEL_YAML.read_text(encoding="utf-8"))["status"] == "provisional"


def test_build_repeated_gives_the_same_final_panel(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    fill_sheet(a)
    fill_sheet(b)
    assert gbd.main(["reconcile"]) == 0
    panel_before, table_before = gbd.PANEL_YAML.read_bytes(), gbd.AMENDMENT_TABLE.read_bytes()
    assert gbd.main(["build", "--status", "final"]) == 0
    assert gbd.PANEL_YAML.read_bytes() == panel_before and gbd.AMENDMENT_TABLE.read_bytes() == table_before
    # ... but not after the ranking or the mapping table was touched
    text = gbd.MAPPING.read_text(encoding="utf-8")
    gbd.MAPPING.write_text(text.replace("23,2022,Alzheimer", "19,2022,Alzheimer", 1), encoding="utf-8")
    assert "19,2022,Alzheimer" in gbd.MAPPING.read_text(encoding="utf-8")
    assert gbd.main(["build", "--status", "final"]) == 2
    assert "does not match the ranking" in capsys.readouterr().err
    assert gbd.PANEL_YAML.read_bytes() == panel_before
    gbd.MAPPING.write_text(text, encoding="utf-8")
    gbd.EXPORT_RECORD.unlink()
    assert gbd.main(["build", "--status", "final"]) == 2
    assert "needs the ranking" in capsys.readouterr().err


def test_fewer_than_twenty_mapped_causes_stop_with_instructions(gbd, work, capsys):
    a, b = prepare_sheets(gbd, work)
    many = {name: (RULE_D, "") for name, _ in LEVEL3[8:24]}
    fill_sheet(a, many)
    fill_sheet(b, many)
    assert gbd.main(["reconcile"]) == 2
    err = capsys.readouterr().err
    assert "the panel needs 20" in err and "--carry-from" in err and not gbd.PANEL_YAML.exists()


def test_build_rebuilds_the_committed_fixture_panel(gbd, tmp_path):
    """tests/fixtures/gbd_panel_fixture.yaml is the output of build for tests/fixtures/gbd_mapping_fixture.csv."""
    mapping = FIXTURES / "gbd_mapping_fixture.csv"
    rc = gbd.build_panel(G.read_csv_dicts(mapping), "final", mapping_text=mapping.read_text(encoding="utf-8"), mapping_path=mapping,
                         panel_yaml=tmp_path / "x" / "p.yaml", table_md=tmp_path / "t.md", test_fixture=True)
    assert rc == 0
    built = yaml.safe_load((tmp_path / "x" / "p.yaml").read_text(encoding="utf-8"))
    committed = yaml.safe_load(FIXTURE_PANEL.read_text(encoding="utf-8"))
    assert built["members"] == committed["members"] and built["status"] == committed["status"] == "final"
    assert committed["test_fixture"] is True
    assert len(committed["members"]) == 20 and committed["members"][-1]["in_trajectory"] is False
    assert G.table_rows_in_text(FIXTURE_AMENDMENT.read_text(encoding="utf-8")) == [G.table_row(m) for m in committed["members"]]


# ---- shared helpers --------------------------------------------------------------------------------------

def test_text_hash_does_not_depend_on_line_endings(gbd, tmp_path):
    lf, crlf = tmp_path / "lf.csv", tmp_path / "crlf.csv"
    lf.write_bytes(b"rank,cause\n1,Stroke\n")
    crlf.write_bytes(b"rank,cause\r\n1,Stroke\r\n")
    assert gbd.sha256_text(lf) == gbd.sha256_text(crlf) == gbd.sha256_file(lf)
    assert gbd.sha256_file(lf) != gbd.sha256_file(crlf)
    gbd.write_text_lf(tmp_path / "out.md", "a\nb\n")
    assert (tmp_path / "out.md").read_bytes() == b"a\nb\n"


def test_table_row_and_key_helpers():
    member = {"rank": 7, "cause": "A | B\ncause", "expression": '"Asthma"[Mesh]', "mesh_ui": "D001249", "date_introduced": "1966-01-01",
              "scope_relation": "exact", "inclusion_note": "line one\nline two | three", "in_trajectory": False, "shared_with": None}
    row = G.table_row(member)
    assert row == '| 7 | A / B cause | "Asthma"[Mesh] | D001249 | 1966 | exact | line one line two / three | queried; 2024 and pooled windows only |'
    assert G.table_rows_in_text("text\n" + row + "\n| 1 | not | a | panel | row |\n| DB | use | template | fields |\n") == [row]
    assert G.gbd_key("Alzheimer's disease and other dementias") == "GBD_ALZHEIMERS_DISEASE_AND_OTHER_DEMENTIAS"
    assert G.gbd_key("HIV/AIDS") == "GBD_HIV_AIDS" and G.norm_name("Diarrhoeal diseases") == G.norm_name("diarrheal  Diseases")
    assert G.rank_of("7.0") == 7
    with pytest.raises(G.GbdPanelError):
        G.rank_of("2 (?)")


def test_final_panel_helper_is_consistent(tmp_path):
    """The end-to-end helper used by the identity-check tests produces a panel that the loader accepts as final."""
    made = final_panel(tmp_path)
    cat = load_catalogue(gbd_panel_path=made["panel"])
    assert cat.gbd_is_final and not cat.gbd_is_test_fixture
    with pytest.raises(CatalogueError, match="must name the two raters"):
        broken = tmp_path / "broken.yaml"
        text = made["panel"].read_text(encoding="utf-8")
        assert "raters: 'IE and WP'" in text
        broken.write_text(text.replace("raters: 'IE and WP'", "raters: ''"), encoding="utf-8")
        load_catalogue(gbd_panel_path=broken)


# ---- the generated text of Section 5.7 in the amendment ------------------------------------------------

def test_amendment_block_check_compares_the_generated_text_as_one_block(tmp_path):
    (tmp_path / "data" / "gbd").mkdir(parents=True)
    table = tmp_path / G.AMENDMENT_TABLE_FILE
    h = ["a" * 64, "b" * 64, "c" * 64, "d" * 64]
    volume = G.volume_line(2292, 4326)
    assert volume == "Catalogue volume with this panel: 2,292 additional count requests, 4,326 per run."
    block = "5.7 GBD panel (final).\nText above.\n\n| 1 | x |\n\nExport: 2026-10-08, by A. Person, " + " ".join(h) + "\n\n" + volume
    table.write_text("\n<!-- generated -->\n\n" + block + "\n", encoding="utf-8")
    panel = {"source": {"mapping_table_sha256": h[1], "ranking_sha256": h[2], "raw_export_sha256": h[3]}}

    def check(text, sha=h[0], required=("| 1 | x |", volume)):
        return G.amendment_block_check(panel, sha, tmp_path, text, required_lines=required)

    assert G.amendment_block(tmp_path) == block                      # without the comment line and the blank lines around it
    assert check("before\n\n" + block + "\n\nafter") == (True, True, "as generated")
    assert check("before\r\n" + block.replace("\n", "  \r\n") + "\r\nafter")[:2] == (True, True)      # line endings, blanks at line ends
    paste_again = "paste the text of data/gbd/gbd_panel_amendment_table.md into the amendment again, unchanged"
    # a difference of blanks only: the message names the place and shows both sides
    strict, lenient, finding = check(block.replace("Text above.", "Text  above."))
    assert (strict, lenient) == (False, True) and finding.endswith(paste_again)
    assert "first difference in line 2 of the generated text, at character 6" in finding and '"5.7 GBD panel (final). / Text above.' in finding
    # a blank line removed
    strict, lenient, finding = check(block.replace("\n\nExport:", "\nExport:"))
    assert (strict, lenient) == (False, True) and "first difference in line 5 of the generated text, at character 1" in finding
    # a character changed far into a long line: the place is found, not only the beginning of the line
    strict, lenient, finding = check(block.replace(h[3], h[3][:-1] + "e"))
    assert (strict, lenient) == (False, False) and "ddddddddd" in finding and "dddde" in finding
    # the text is not there at all
    assert check("nothing of it")[2].startswith("the generated text does not begin anywhere in the amendment")

    # the deposited text was not written for the files as they are now
    assert "does not state the SHA-256" in check(block, sha="f" * 64)[2]
    panel["source"]["raters"] = "IE and WP"
    assert "does not name the raters of the panel file (IE and WP)" in check(block)[2]
    with_raters = block.replace("Text above.", "Mapped by the rule of 5.3 by IE and WP. Text above.")
    table.write_text(with_raters + "\n", encoding="utf-8")
    assert check(with_raters)[0] and not check(with_raters.replace("IE and WP", "XX and YY"))[0]
    del panel["source"]["raters"]
    table.write_text("<!-- generated -->\n" + block + "\n", encoding="utf-8")
    assert "does not hold the table of the panel file and the catalogue volume" in check(block, required=("| 1 | y |", volume))[2]
    assert "does not hold the table of the panel file and the catalogue volume" in check(block, required=("| 1 | x |", G.volume_line(1, 2)))[2]
    record = tmp_path / "data" / "gbd" / "gbd_export_record.json"
    record.write_text(json.dumps({"export_date": "2026-10-08", "exported_by": "A. Person", "permalink": "", "cause_hierarchy": {"sha256": h[2]}}), encoding="utf-8")
    assert check(block)[0]
    record.write_text(json.dumps({"export_date": "2026-10-08", "exported_by": "Somebody Else"}), encoding="utf-8")
    assert "does not state the export date, person, permalink and hierarchy hash" in check(block)[2]
    record.unlink()
    # a byte order mark in the deposited text is no difference; a file without text is reported as such
    table.write_bytes(b"\xef\xbb\xbf" + ("<!-- generated -->\n" + block + "\n").encode("utf-8"))
    assert check(block)[0]
    table.write_text("<!-- generated -->\n\n", encoding="utf-8")
    assert "not found or without text" in check(block)[2]
    table.unlink()
    assert G.amendment_block(tmp_path) == "" and "not found or without text" in check(block)[2]
